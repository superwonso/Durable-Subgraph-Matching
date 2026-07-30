import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from create_selectivity_ablation_excel import (
    EXPECTED_SHEETS,
    build_workbook,
    load_inputs,
    validate_workbook,
)


class SelectivityAblationWorkbookTests(unittest.TestCase):
    def make_inputs(
        self,
        root: Path,
        *,
        statistics_off_zero: bool = False,
        correctness_status: str = "pass",
        candidate_schema_mismatch: bool = False,
    ) -> Path:
        input_dir = root / "evaluation"
        input_dir.mkdir()
        algorithms = {
            "original_perf": {
                "executable": "/tmp/original/td_tree",
                "label": "Original (TD-Tree)",
                "cohort": "performance",
                "flags": ["--duration-mode", "consecutive", "--threads", "32"],
            },
            "ours_tdtree_perf": {
                "executable": "/tmp/ours/td_tree",
                "cohort": "performance",
                "flags": [
                    "--selectivity",
                    "tdtree",
                    "--prefilter",
                    "on",
                    "--threads",
                    "32",
                ],
            },
            "gopher_cpu_perf": {
                "executable": "/tmp/ours_gopher/td_tree_gopher",
                "cohort": "performance",
                "flags": [
                    "--statistics",
                    "off",
                    "--selectivity",
                    "paper",
                    "--threads",
                    "32",
                ],
            },
            "gopher_gpu_perf": {
                "executable": "/tmp/ours_gopher_gpu/td_tree_gopher_gpu",
                "cohort": "performance",
                "flags": [
                    "--backend",
                    "gpu",
                    "--statistics",
                    "off",
                    "--gpu-variant",
                    "resident",
                    "--gpu-memory-mb",
                    "64",
                ],
            },
            "gopher_prune_on_diag": {
                "executable": "/tmp/ours_gopher/td_tree_gopher",
                "cohort": "diagnostic",
                "flags": [
                    "--statistics",
                    "on",
                    "--early-prune",
                    "on",
                    "--threads",
                    "1",
                ],
            },
            "gopher_prune_off_diag": {
                "executable": "/tmp/ours_gopher/td_tree_gopher",
                "cohort": "diagnostic",
                "flags": [
                    "--statistics",
                    "on",
                    "--early-prune",
                    "off",
                    "--threads",
                    "1",
                ],
            },
            "gopher_gpu_diag": {
                "executable": "/tmp/ours_gopher_gpu/td_tree_gopher_gpu",
                "cohort": "diagnostic",
                "flags": [
                    "--backend",
                    "gpu",
                    "--statistics",
                    "on",
                    "--gpu-variant",
                    "resident",
                ],
            },
        }
        performance_algorithms = [
            "original_perf",
            "ours_tdtree_perf",
            "gopher_cpu_perf",
            "gopher_gpu_perf",
        ]
        diagnostic_algorithms = [
            "gopher_prune_on_diag",
            "gopher_prune_off_diag",
            "gopher_gpu_diag",
        ]
        cases = [
            {
                "id": "selectivity_case",
                "suite": "selectivity_performance",
                "dataset_id": "fixture",
                "query_shape": "path3",
                "k": 3,
                "repeats": 2,
                "expected_match_count": 10,
                "algorithms": performance_algorithms,
            },
            {
                "id": "prune_case",
                "suite": "pruning_diagnostic",
                "dataset_id": "fixture",
                "query_shape": "path3",
                "k": 4,
                "repeats": 2,
                "expected_match_count": 7,
                "algorithms": diagnostic_algorithms,
            },
        ]
        manifest = {
            "schema_version": 1,
            "title": "fixture",
            "repeats": 2,
            "methodology": {"primary_metric": "median wall_s"},
            "algorithms": algorithms,
            "cases": cases,
            "comparisons": [
                {
                    "id": "ours_vs_original",
                    "case_id": "selectivity_case",
                    "baseline_algorithm": "original_perf",
                    "variant_algorithms": ["ours_tdtree_perf"],
                    "cohort": "performance",
                    "metric": "wall_s",
                },
                {
                    "id": "cpu_gpu_side_by_side",
                    "case_id": "selectivity_case",
                    "baseline_algorithm": "gopher_cpu_perf",
                    "variant_algorithms": ["gopher_gpu_perf"],
                    "cohort": "performance",
                    "metric": "wall_s",
                },
                {
                    "id": "prune_candidates",
                    "case_id": "prune_case",
                    "baseline_algorithm": "gopher_prune_off_diag",
                    "variant_algorithms": ["gopher_prune_on_diag"],
                    "cohort": "diagnostic",
                    "candidate_metric": "raw_extension_candidates",
                },
                {
                    "id": "cpu_gpu_candidates",
                    "case_id": "prune_case",
                    "baseline_algorithm": "gopher_prune_on_diag",
                    "variant_algorithms": ["gopher_gpu_diag"],
                    "cohort": "diagnostic",
                    "candidate_metric": "raw_extension_candidates",
                },
            ],
        }
        (input_dir / "manifest.used.json").write_text(
            json.dumps(manifest), encoding="utf-8"
        )
        (input_dir / "environment.json").write_text(
            json.dumps({"cpu": "fixture", "gpu": "fixture"}), encoding="utf-8"
        )

        fields = [
            "case_id",
            "algorithm",
            "repeat",
            "status",
            "wall_s",
            "match_count",
            "backend",
            "threads_used",
            "statistics_enabled",
            "selectivity_config",
            "prefilter_config",
            "early_prune",
            "metric_schema_version",
            "root_candidate_count",
            "raw_extension_candidates",
            "accepted_extensions",
            "dfs_states",
            "completed_matches",
            "pruned_injectivity",
            "pruned_temporal_support",
            "candidate_relation_entries",
            "retained_edges",
            "gpu_variant",
            "gpu_memory_mb_config",
            "gpuKernel_ms",
            "gpuReadyToResult_ms",
            "fallback_used",
            "resource_max_rss_kb",
            "run_signature",
            "command",
        ]
        timing = {
            "original_perf": 2.0,
            "ours_tdtree_perf": 1.0,
            "gopher_cpu_perf": 0.8,
            "gopher_gpu_perf": 0.6,
            "gopher_prune_on_diag": 1.1,
            "gopher_prune_off_diag": 1.5,
            "gopher_gpu_diag": 0.9,
        }
        with (input_dir / "runs.csv").open(
            "w", newline="", encoding="utf-8"
        ) as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            for case in cases:
                for algorithm_id in case["algorithms"]:
                    config = algorithms[algorithm_id]
                    is_gpu = "gpu" in algorithm_id
                    is_diag = config["cohort"] == "diagnostic"
                    stats_off = algorithm_id in {
                        "gopher_cpu_perf",
                        "gopher_gpu_perf",
                    }
                    for repeat in range(2):
                        raw_candidates = (
                            0
                            if statistics_off_zero and stats_off
                            else ""
                            if stats_off
                            else 1000
                            if algorithm_id != "gopher_prune_on_diag"
                            else 600
                        )
                        writer.writerow(
                            {
                                "case_id": case["id"],
                                "algorithm": algorithm_id,
                                "repeat": repeat,
                                "status": "ok",
                                "wall_s": timing[algorithm_id] + repeat * 0.02,
                                "match_count": case["expected_match_count"],
                                "backend": "gpu" if is_gpu else "cpu",
                                "threads_used": 1 if is_gpu or is_diag else 32,
                                "statistics_enabled": (
                                    "false" if stats_off else "true" if is_diag else ""
                                ),
                                "selectivity_config": (
                                    "tdtree"
                                    if algorithm_id == "ours_tdtree_perf"
                                    else "n/a"
                                    if algorithm_id == "original_perf"
                                    else "paper"
                                ),
                                "prefilter_config": (
                                    "on"
                                    if algorithm_id == "ours_tdtree_perf"
                                    else "n/a"
                                ),
                                "early_prune": (
                                    "off"
                                    if algorithm_id == "gopher_prune_off_diag"
                                    else "on"
                                    if "gopher" in algorithm_id
                                    else "n/a"
                                ),
                                "metric_schema_version": (
                                    "logical_extensions_v2"
                                    if candidate_schema_mismatch
                                    and algorithm_id == "gopher_prune_on_diag"
                                    else "logical_extensions_v1"
                                ),
                                "root_candidate_count": (
                                    "" if stats_off else 20
                                ),
                                "raw_extension_candidates": raw_candidates,
                                "accepted_extensions": (
                                    "" if stats_off else 300
                                ),
                                "dfs_states": "" if stats_off else 320,
                                "completed_matches": "" if stats_off else case["expected_match_count"],
                                "pruned_injectivity": "" if stats_off else 40,
                                "pruned_temporal_support": "" if stats_off else 20,
                                "candidate_relation_entries": 1500,
                                "retained_edges": 500,
                                "gpu_variant": "resident" if is_gpu else "",
                                "gpu_memory_mb_config": 64 if is_gpu else "",
                                "gpuKernel_ms": 10 if is_gpu else "",
                                "gpuReadyToResult_ms": 50 if is_gpu else "",
                                "fallback_used": "false",
                                "resource_max_rss_kb": 1000,
                                "run_signature": f"{case['id']}-{algorithm_id}-{repeat}",
                                "command": f"/tmp/{algorithm_id}",
                            }
                        )

        with (input_dir / "summary.csv").open(
            "w", newline="", encoding="utf-8"
        ) as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=["case_id", "algorithm", "successful_runs", "wall_s_median"],
            )
            writer.writeheader()
            for case in cases:
                for algorithm_id in case["algorithms"]:
                    writer.writerow(
                        {
                            "case_id": case["id"],
                            "algorithm": algorithm_id,
                            "successful_runs": 2,
                            "wall_s_median": timing[algorithm_id] + 0.01,
                        }
                    )

        with (input_dir / "correctness.csv").open(
            "w", newline="", encoding="utf-8"
        ) as handle:
            writer = csv.DictWriter(
                handle, fieldnames=["case_id", "semantic_group", "status"]
            )
            writer.writeheader()
            for case in cases:
                writer.writerow(
                    {
                        "case_id": case["id"],
                        "semantic_group": "consecutive",
                        "status": correctness_status,
                    }
                )

        with (input_dir / "candidate_rows.csv").open(
            "w", newline="", encoding="utf-8"
        ) as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=["case_id", "algorithm", "query_vertex", "relation_entries"],
            )
            writer.writeheader()
            writer.writerow(
                {
                    "case_id": "prune_case",
                    "algorithm": "gopher_prune_on_diag",
                    "query_vertex": 1,
                    "relation_entries": 100,
                }
            )
        with (input_dir / "selectivity_rows.csv").open(
            "w", newline="", encoding="utf-8"
        ) as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=[
                    "case_id",
                    "algorithm",
                    "query_vertex",
                    "score",
                    "selected_root",
                ],
            )
            writer.writeheader()
            writer.writerow(
                {
                    "case_id": "selectivity_case",
                    "algorithm": "ours_tdtree_perf",
                    "query_vertex": 0,
                    "score": 0.1,
                    "selected_root": "true",
                }
            )
        return input_dir

    def test_build_and_validate_strict_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = self.make_inputs(root)
            output = root / "selectivity.xlsx"
            receipt = build_workbook(input_dir, output)
            self.assertEqual(receipt["validation"], "passed")
            self.assertEqual(receipt["official_runs"], 14)
            self.assertEqual(receipt["metric_schema_versions"], ["logical_extensions_v1"])
            self.assertGreaterEqual(receipt["charts"], 3)
            self.assertEqual(validate_workbook(output)["sheets"], 10)

            workbook = load_workbook(output, data_only=True)
            self.assertEqual(workbook.sheetnames, EXPECTED_SHEETS)
            controlled = workbook["Controlled Ablation"]
            header_row = next(
                row
                for row in range(1, controlled.max_row + 1)
                if controlled.cell(row, 1).value == "Comparison ID"
            )
            headers = {
                controlled.cell(header_row, column).value: column
                for column in range(1, controlled.max_column + 1)
            }
            result = {}
            for row in range(header_row + 1, controlled.max_row + 1):
                values = {
                    name: controlled.cell(row, column).value
                    for name, column in headers.items()
                }
                result[
                    (
                        controlled.cell(row, 1).value,
                        values["Metric"],
                    )
                ] = values
            self.assertTrue(
                result[("ours_vs_original", "wall_s")]["Ratio eligible"]
            )
            self.assertAlmostEqual(
                result[("ours_vs_original", "wall_s")]["Baseline / variant"],
                2.0,
                places=1,
            )
            self.assertFalse(
                result[("cpu_gpu_side_by_side", "wall_s")]["Ratio eligible"]
            )
            self.assertIsNone(
                result[("cpu_gpu_side_by_side", "wall_s")][
                    "Baseline / variant"
                ]
            )
            self.assertIn(
                "CPU/GPU",
                result[("cpu_gpu_side_by_side", "wall_s")][
                    "Suppression reason"
                ],
            )
            self.assertTrue(
                result[("prune_candidates", "raw_extension_candidates")][
                    "Ratio eligible"
                ]
            )
            self.assertFalse(
                result[("cpu_gpu_candidates", "raw_extension_candidates")][
                    "Ratio eligible"
                ]
            )

    def test_statistics_off_zero_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            input_dir = self.make_inputs(
                Path(temporary), statistics_off_zero=True
            )
            with self.assertRaisesRegex(ValueError, "must serialize.*null"):
                load_inputs(input_dir)

    def test_nonpass_correctness_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            input_dir = self.make_inputs(
                Path(temporary), correctness_status="fail"
            )
            with self.assertRaisesRegex(ValueError, "correctness did not pass"):
                load_inputs(input_dir)

    def test_candidate_ratio_requires_equal_metric_schema(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = self.make_inputs(
                root, candidate_schema_mismatch=True
            )
            output = root / "schema_mismatch.xlsx"
            build_workbook(input_dir, output)
            workbook = load_workbook(output, data_only=True)
            ws = workbook["Controlled Ablation"]
            header_row = next(
                row
                for row in range(1, ws.max_row + 1)
                if ws.cell(row, 1).value == "Comparison ID"
            )
            headers = {
                ws.cell(header_row, column).value: column
                for column in range(1, ws.max_column + 1)
            }
            row = next(
                row
                for row in range(header_row + 1, ws.max_row + 1)
                if ws.cell(row, 1).value == "prune_candidates"
                and ws.cell(row, headers["Metric"]).value
                == "raw_extension_candidates"
            )
            self.assertFalse(ws.cell(row, headers["Ratio eligible"]).value)
            self.assertIsNone(
                ws.cell(row, headers["Baseline / variant"]).value
            )
            self.assertIn(
                "schema",
                ws.cell(row, headers["Suppression reason"]).value,
            )


if __name__ == "__main__":
    unittest.main()
