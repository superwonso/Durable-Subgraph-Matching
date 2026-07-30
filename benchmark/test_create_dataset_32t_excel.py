import csv
import json
import statistics
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from create_dataset_32t_excel import (
    ALGORITHM_ORDER,
    BASELINE_ALGORITHM,
    DATASET_ORDER,
    EXPECTED_SHEETS,
    EXPECTED_SHEETS_WITH_FILTER,
    FILTER_RAW_COLUMNS,
    STABLE_LABEL_MODE,
    build_workbook,
    gather,
    validate_workbook,
)


class Dataset32TWorkbookTests(unittest.TestCase):
    def make_inputs(
        self,
        root: Path,
        *,
        zero_dataset: str | None = None,
        disagreement: bool = False,
        wrong_data_path: bool = False,
        wrong_input_variant: bool = False,
        wrong_label_mode: bool = False,
    ) -> Path:
        input_dir = root / "results"
        input_dir.mkdir()
        data_dir = root / "data"
        data_dir.mkdir()
        algorithms = {}
        for algorithm in ALGORITHM_ORDER:
            algorithms[algorithm] = {
                "executable": f"/tmp/{algorithm}",
                "semantic_group": "consecutive",
                "role": algorithm,
                "threads": 1 if algorithm == "ours_gopher_gpu" else 32,
                "flags": [
                    "--count-only",
                    "--label-mode",
                    STABLE_LABEL_MODE,
                ],
            }
        cases = []
        for dataset in DATASET_ORDER:
            raw_path = data_dir / f"{dataset}.txt"
            filtered_path = data_dir / f"{dataset}_filtered.txt"
            raw_path.write_text("1 2 1\n", encoding="utf-8")
            filtered_path.write_text("1 2 1\n1 2 2\n", encoding="utf-8")
            cases.append(
                {
                    "id": f"case_{dataset}",
                    "dataset_id": dataset,
                    "data": str(filtered_path),
                    "data_by_algorithm": {
                        BASELINE_ALGORITHM: str(raw_path),
                    },
                    "input_variant": "supplied_prefiltered_proposed",
                    "input_variant_by_algorithm": {
                        BASELINE_ALGORITHM: "raw_unique_occurrences_baseline",
                    },
                    "query": f"/tmp/{dataset}.qry",
                    "k": 3,
                    "repeats": 2,
                    "algorithms": ALGORITHM_ORDER,
                }
            )
        manifest = {
            "schema_version": 1,
            "repeats": 2,
            "algorithms": algorithms,
            "cases": cases,
            "methodology": {"primary_metric": "median wall"},
        }
        (input_dir / "manifest.used.json").write_text(
            json.dumps(manifest), encoding="utf-8"
        )
        (input_dir / "environment.json").write_text(
            json.dumps({"cpu": {"logical_threads": 32}}), encoding="utf-8"
        )

        run_fields = [
            "case_id",
            "dataset_id",
            "algorithm",
            "input_variant",
            "data",
            "repeat",
            "status",
            "wall_s",
            "match_count",
            "threads_requested",
            "threads_used",
            "workers_with_tasks",
            "parallel_workers_with_tasks",
            "cpu_affinity",
            "fallback_used",
            "gpu_variant",
            "gpuKernel_ms",
            "gpuReadyToResult_ms",
            "resource_max_rss_kb",
            "executable_sha256",
            "data_sha256",
            "query_sha256",
            "run_signature",
            "label_mode_config",
            "input_occurrences",
            "active_vertices",
            "command",
        ]
        with (input_dir / "runs.csv").open(
            "w", newline="", encoding="utf-8"
        ) as handle:
            writer = csv.DictWriter(handle, fieldnames=run_fields)
            writer.writeheader()
            for dataset_index, dataset in enumerate(DATASET_ORDER):
                base_count = 0 if dataset == zero_dataset else (dataset_index + 1) * 10
                for algorithm_index, algorithm in enumerate(ALGORITHM_ORDER):
                    is_baseline = algorithm == BASELINE_ALGORITHM
                    recorded_data = (
                        data_dir / f"{dataset}.txt"
                        if is_baseline
                        else data_dir / f"{dataset}_filtered.txt"
                    )
                    input_variant = (
                        "raw_unique_occurrences_baseline"
                        if is_baseline
                        else "supplied_prefiltered_proposed"
                    )
                    count = (
                        base_count + 1
                        if disagreement
                        and dataset == DATASET_ORDER[0]
                        and algorithm == ALGORITHM_ORDER[-1]
                        else base_count
                    )
                    for repeat in range(2):
                        writer.writerow(
                            {
                                "case_id": f"case_{dataset}",
                                "dataset_id": dataset,
                                "algorithm": algorithm,
                                "input_variant": (
                                    "wrong_variant"
                                    if wrong_input_variant
                                    and dataset == DATASET_ORDER[0]
                                    and algorithm == BASELINE_ALGORITHM
                                    else input_variant
                                ),
                                "data": (
                                    str(data_dir / "unexpected.txt")
                                    if wrong_data_path
                                    and dataset == DATASET_ORDER[0]
                                    and algorithm == BASELINE_ALGORITHM
                                    else str(recorded_data)
                                ),
                                "repeat": repeat,
                                "status": "ok",
                                "wall_s": 1.0 + dataset_index + algorithm_index * 0.1
                                + repeat * 0.01,
                                "match_count": count,
                                "threads_requested": (
                                    1 if algorithm == "ours_gopher_gpu" else 32
                                ),
                                "threads_used": (
                                    1 if algorithm == "ours_gopher_gpu" else 32
                                ),
                                "workers_with_tasks": (
                                    "" if algorithm == "ours_gopher_gpu" else 8
                                ),
                                "parallel_workers_with_tasks": "",
                                "cpu_affinity": (
                                    "" if algorithm == "ours_gopher_gpu" else "0-31"
                                ),
                                "fallback_used": "false",
                                "gpu_variant": (
                                    "resident"
                                    if algorithm == "ours_gopher_gpu"
                                    else ""
                                ),
                                "gpuKernel_ms": (
                                    10 if algorithm == "ours_gopher_gpu" else ""
                                ),
                                "gpuReadyToResult_ms": (
                                    100 if algorithm == "ours_gopher_gpu" else ""
                                ),
                                "resource_max_rss_kb": 1000,
                                "executable_sha256": algorithm * 2,
                                "data_sha256": (
                                    f"raw-{dataset}"
                                    if is_baseline
                                    else f"filtered-{dataset}"
                                ),
                                "query_sha256": dataset * 3,
                                "run_signature": f"{dataset}-{algorithm}-{repeat}",
                                "label_mode_config": (
                                    "sequential-rng"
                                    if wrong_label_mode
                                    and dataset == DATASET_ORDER[0]
                                    and algorithm == BASELINE_ALGORITHM
                                    else STABLE_LABEL_MODE
                                ),
                                "input_occurrences": (dataset_index + 1) * 1000,
                                "active_vertices": (dataset_index + 1) * 100,
                                "command": f"/tmp/{algorithm} data query 3",
                            }
                        )

        correctness_fields = [
            "case_id",
            "semantic_group",
            "status",
        ]
        with (input_dir / "correctness.csv").open(
            "w", newline="", encoding="utf-8"
        ) as handle:
            writer = csv.DictWriter(handle, fieldnames=correctness_fields)
            writer.writeheader()
            for dataset in DATASET_ORDER:
                writer.writerow(
                    {
                        "case_id": f"case_{dataset}",
                        "semantic_group": "consecutive",
                        "status": "pass",
                    }
                )
        return input_dir

    def make_filter_inputs(
        self,
        root: Path,
        input_dir: Path,
        *,
        incomplete: bool = False,
        canonical_failure: bool = False,
        summary_mismatch: bool = False,
        wrong_raw_path: bool = False,
    ) -> Path:
        filter_dir = root / "filter_results"
        filter_dir.mkdir()
        manifest = json.loads(
            (input_dir / "manifest.used.json").read_text(encoding="utf-8")
        )
        cases = {
            case["dataset_id"]: case for case in manifest["cases"]
        }
        repetitions = int(manifest["repeats"])
        source_hash = "a" * 64
        executable_hash = "b" * 64
        rows: list[dict[str, object]] = []
        canonical_records: dict[str, dict[str, object]] = {}

        for repeat in range(1, repetitions + 1):
            rotation = (repeat - 1) % len(DATASET_ORDER)
            order = DATASET_ORDER[rotation:] + DATASET_ORDER[:rotation]
            for position, dataset in enumerate(order, start=1):
                dataset_index = DATASET_ORDER.index(dataset)
                case = cases[dataset]
                raw_path = Path(
                    case["data_by_algorithm"][BASELINE_ALGORITHM]
                )
                if wrong_raw_path and dataset == DATASET_ORDER[0]:
                    raw_path = root / "wrong-raw.txt"
                filtered_path = Path(case["data"])
                wall_s = 1.0 + dataset_index + repeat * 0.1
                filter_core_ms = 10 + dataset_index
                serialize_loop_ms = 5 + dataset_index
                remainder_ms = (
                    wall_s * 1000.0
                    - filter_core_ms
                    - serialize_loop_ms
                )
                generated_rows = (dataset_index + 1) * 1000
                canonical_hash = f"{dataset_index + 1:064x}"
                canonical_status = (
                    "fail"
                    if canonical_failure
                    and dataset == DATASET_ORDER[0]
                    else "pass"
                )
                row = {
                    "dataset": dataset,
                    "repeat": repeat,
                    "execution_position": position,
                    "status": "ok",
                    "exit_code": 0,
                    "wall_s": f"{wall_s:.9f}",
                    "filter_core_ms": filter_core_ms,
                    "serialize_loop_ms": serialize_loop_ms,
                    "unattributed_read_group_close_ms": (
                        f"{remainder_ms:.6f}"
                    ),
                    "raw_unique_edges": 100 + dataset_index,
                    "retained_unique_edges": 10 + dataset_index,
                    "generated_rows": generated_rows,
                    "generated_bytes": generated_rows * 10,
                    "generated_sha256": f"{dataset_index + 11:064x}",
                    "canonical_validation": canonical_status,
                    "canonical_generated_sha256": (
                        canonical_hash if repeat == 1 else ""
                    ),
                    "canonical_supplied_sha256": (
                        canonical_hash if repeat == 1 else ""
                    ),
                    "max_rss_kb": 100_000 + dataset_index * 1000,
                    "major_faults": 0,
                    "minor_faults": 100,
                    "voluntary_context_switches": 1,
                    "involuntary_context_switches": 2,
                    "cpu_affinity": "0",
                    "raw_path": str(raw_path),
                    "supplied_filtered_path": str(filtered_path),
                    "source_sha256": source_hash,
                    "executable_sha256": executable_hash,
                    "started_at_utc": (
                        f"2026-07-29T00:00:0{repeat}+00:00"
                    ),
                    "ended_at_utc": (
                        f"2026-07-29T00:00:1{repeat}+00:00"
                    ),
                }
                rows.append(row)
                canonical_records[dataset] = {
                    "status": canonical_status,
                    "generated_sha256": canonical_hash,
                    "supplied_sha256": canonical_hash,
                    "generated_rows": generated_rows,
                    "generated_bytes": generated_rows * 10,
                }
        if incomplete:
            rows.pop()

        with (filter_dir / "runs.csv").open(
            "w", newline="", encoding="utf-8"
        ) as handle:
            writer = csv.DictWriter(handle, fieldnames=FILTER_RAW_COLUMNS)
            writer.writeheader()
            writer.writerows(rows)

        summary_rows: list[dict[str, object]] = []
        numeric_fields = (
            "wall_s",
            "filter_core_ms",
            "serialize_loop_ms",
            "unattributed_read_group_close_ms",
            "max_rss_kb",
        )
        for dataset in DATASET_ORDER:
            group = [row for row in rows if row["dataset"] == dataset]
            if not group:
                continue
            summary_row: dict[str, object] = {
                "dataset": dataset,
                "repetitions": len(group),
                "canonical_validation": group[0][
                    "canonical_validation"
                ],
                "raw_unique_edges": group[0]["raw_unique_edges"],
                "retained_unique_edges": group[0][
                    "retained_unique_edges"
                ],
                "generated_rows": group[0]["generated_rows"],
                "generated_bytes": group[0]["generated_bytes"],
            }
            for field in numeric_fields:
                values = [float(row[field]) for row in group]
                summary_row[f"{field}_median"] = statistics.median(
                    values
                )
                summary_row[f"{field}_min"] = min(values)
                summary_row[f"{field}_max"] = max(values)
            summary_rows.append(summary_row)
        if summary_mismatch:
            summary_rows[0]["wall_s_median"] = (
                float(summary_rows[0]["wall_s_median"]) + 1.0
            )
        with (filter_dir / "summary.csv").open(
            "w", newline="", encoding="utf-8"
        ) as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=list(summary_rows[0]),
            )
            writer.writeheader()
            writer.writerows(summary_rows)

        validation_status = (
            "fail" if canonical_failure or incomplete else "pass"
        )
        validation = {
            "status": validation_status,
            "datasets": canonical_records,
            "observed_runs": len(rows),
            "expected_runs": len(DATASET_ORDER) * repetitions,
            "non_ok_runs": 0,
        }
        (filter_dir / "validation.json").write_text(
            json.dumps(validation), encoding="utf-8"
        )
        environment = {
            "captured_at_utc": "2026-07-29T00:00:00+00:00",
            "source": "/tmp/effective_filtering_with_time_cond.cpp",
            "source_sha256": source_hash,
            "compile_command": ["g++", "-O3"],
            "executable_sha256": executable_hash,
            "repetitions": repetitions,
            "cpu_affinity": "0",
            "page_cache_policy": "OS page cache not flushed",
            "timing_boundaries": {
                "wall_s": "full process",
                "filter_core_ms": "narrow core",
                "serialize_loop_ms": "narrow output loop",
            },
        }
        (filter_dir / "environment.json").write_text(
            json.dumps(environment), encoding="utf-8"
        )
        return filter_dir

    def test_build_and_validate(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = self.make_inputs(root)
            output = root / "report.xlsx"
            receipt = build_workbook(input_dir, output)
            self.assertEqual(receipt["validation"], "passed")
            self.assertEqual(receipt["comparison_rows"], 16)
            self.assertEqual(receipt["input_protocol_rows"], 16)
            self.assertEqual(receipt["zero_match_datasets"], 0)
            self.assertEqual(receipt["filter_materialization_rows"], 0)
            self.assertEqual(receipt["filter_amortization_rows"], 0)
            self.assertEqual(receipt["filter_raw_runs"], 0)
            self.assertEqual(validate_workbook(output)["charts"], 6)

            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook.sheetnames, EXPECTED_SHEETS)
            overview = workbook["Overview"]
            overview_values = {
                overview.cell(row, 1).value: overview.cell(row, 2).value
                for row in range(1, overview.max_row + 1)
            }
            self.assertEqual(
                overview_values["Input protocol"],
                "TD-Tree=raw; proposed=precomputed filtered",
            )
            self.assertEqual(
                overview_values["Label mode"], STABLE_LABEL_MODE
            )
            self.assertEqual(
                overview_values["Timing boundary"],
                "Precomputed-filter online comparison",
            )

            raw_ws = workbook["Raw Runs"]
            raw_header = next(
                row
                for row in range(1, raw_ws.max_row + 1)
                if raw_ws.cell(row, 1).value == "case_id"
            )
            raw_headers = {
                raw_ws.cell(raw_header, column).value: column
                for column in range(1, raw_ws.max_column + 1)
            }
            baseline_row = next(
                row
                for row in range(raw_header + 1, raw_ws.max_row + 1)
                if raw_ws.cell(
                    row, raw_headers["algorithm"]
                ).value == BASELINE_ALGORITHM
            )
            proposed_row = next(
                row
                for row in range(raw_header + 1, raw_ws.max_row + 1)
                if raw_ws.cell(
                    row, raw_headers["algorithm"]
                ).value == "ours_cpu_32t"
            )
            self.assertNotIn(
                "_filtered",
                raw_ws.cell(baseline_row, raw_headers["data"]).value,
            )
            self.assertIn(
                "_filtered",
                raw_ws.cell(proposed_row, raw_headers["data"]).value,
            )
            self.assertEqual(
                raw_ws.cell(
                    baseline_row, raw_headers["label_mode_config"]
                ).value,
                STABLE_LABEL_MODE,
            )

            configuration = workbook["Configuration"]
            protocol_header = next(
                row
                for row in range(1, configuration.max_row + 1)
                if configuration.cell(row, 1).value == "Dataset ID"
            )
            protocol_rows = [
                row
                for row in configuration.iter_rows(
                    min_row=protocol_header + 1,
                    max_row=configuration.max_row,
                    values_only=True,
                )
                if row[0]
            ]
            self.assertEqual(len(protocol_rows), 16)
            self.assertTrue(all(row[12] is True for row in protocol_rows))
            self.assertEqual(
                {row[3] for row in protocol_rows},
                {"raw baseline", "precomputed filtered proposed"},
            )

    def test_build_with_filter_materialization_and_validate(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = self.make_inputs(root)
            filter_dir = self.make_filter_inputs(root, input_dir)
            output = root / "report-with-filter.xlsx"
            receipt = build_workbook(input_dir, output, filter_dir)

            self.assertEqual(receipt["validation"], "passed")
            self.assertEqual(receipt["filter_materialization_rows"], 4)
            self.assertEqual(receipt["filter_amortization_rows"], 12)
            self.assertEqual(receipt["filter_raw_runs"], 8)
            self.assertEqual(validate_workbook(output)["charts"], 6)

            workbook = load_workbook(output, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                EXPECTED_SHEETS_WITH_FILTER,
            )
            filter_ws = workbook["Filter Materialization"]
            summary_header = next(
                row
                for row in range(1, filter_ws.max_row + 1)
                if filter_ws.cell(row, 1).value == "Dataset"
                and filter_ws.cell(row, 2).value == "Runs"
            )
            summary_headers = {
                filter_ws.cell(summary_header, column).value
                for column in range(1, filter_ws.max_column + 1)
            }
            self.assertIn(
                "Unattributed remainder median (ms)",
                summary_headers,
            )
            self.assertNotIn("Read time median (ms)", summary_headers)
            amortization_header = next(
                row
                for row in range(1, filter_ws.max_row + 1)
                if filter_ws.cell(row, 1).value == "Dataset"
                and filter_ws.cell(row, 2).value == "Method"
            )
            amortization_headers = {
                filter_ws.cell(amortization_header, column).value
                for column in range(1, filter_ws.max_column + 1)
            }
            self.assertIn(
                "10-query amortized per-query (s)",
                amortization_headers,
            )
            self.assertNotIn(
                "5-query amortized per-query (s)",
                amortization_headers,
            )

            overview = workbook["Overview"]
            overview_values = {
                overview.cell(row, 1).value: overview.cell(row, 2).value
                for row in range(1, overview.max_row + 1)
            }
            self.assertEqual(overview_values["Filter validation"], "PASS")
            self.assertEqual(
                overview_values["Filter timing primary"],
                "Full raw-to-filtered wall time",
            )

    def test_incomplete_filter_materialization_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = self.make_inputs(root)
            filter_dir = self.make_filter_inputs(
                root,
                input_dir,
                incomplete=True,
            )
            with self.assertRaisesRegex(
                ValueError,
                "run count mismatch",
            ):
                build_workbook(
                    input_dir,
                    root / "report.xlsx",
                    filter_dir,
                )

    def test_filter_summary_mismatch_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = self.make_inputs(root)
            filter_dir = self.make_filter_inputs(
                root,
                input_dir,
                summary_mismatch=True,
            )
            with self.assertRaisesRegex(
                ValueError,
                "summary disagrees",
            ):
                build_workbook(
                    input_dir,
                    root / "report.xlsx",
                    filter_dir,
                )

    def test_filter_input_path_mismatch_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = self.make_inputs(root)
            filter_dir = self.make_filter_inputs(
                root,
                input_dir,
                wrong_raw_path=True,
            )
            with self.assertRaisesRegex(
                ValueError,
                "raw path does not match",
            ):
                build_workbook(
                    input_dir,
                    root / "report.xlsx",
                    filter_dir,
                )

    def test_filter_canonical_failure_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = self.make_inputs(root)
            filter_dir = self.make_filter_inputs(
                root,
                input_dir,
                canonical_failure=True,
            )
            with self.assertRaisesRegex(
                ValueError,
                "canonical validation did not pass",
            ):
                build_workbook(
                    input_dir,
                    root / "report.xlsx",
                    filter_dir,
                )

    def test_disagreement_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            input_dir = self.make_inputs(Path(temporary), disagreement=True)
            with self.assertRaisesRegex(ValueError, "disagree"):
                gather(input_dir)

    def test_zero_selected_dataset_is_rejected_for_handoff(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = self.make_inputs(
                root, zero_dataset=DATASET_ORDER[0]
            )
            with self.assertRaisesRegex(ValueError, "zero or missing"):
                build_workbook(input_dir, root / "report.xlsx")

    def test_recorded_data_path_must_match_algorithm_manifest_path(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            input_dir = self.make_inputs(
                Path(temporary), wrong_data_path=True
            )
            with self.assertRaisesRegex(
                ValueError, "does not match manifest"
            ):
                gather(input_dir)

    def test_recorded_input_variant_must_match_manifest(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            input_dir = self.make_inputs(
                Path(temporary), wrong_input_variant=True
            )
            with self.assertRaisesRegex(
                ValueError, "recorded input_variant"
            ):
                gather(input_dir)

    def test_stable_id_label_mode_is_required(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            input_dir = self.make_inputs(
                Path(temporary), wrong_label_mode=True
            )
            with self.assertRaisesRegex(
                ValueError, "label_mode_config=stable-id"
            ):
                gather(input_dir)


if __name__ == "__main__":
    unittest.main()
