#!/usr/bin/env python3
"""Create the focused dataset-wise 32-thread/GPU comparison workbook."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import statistics
import tempfile
from collections import defaultdict
from copy import copy
from pathlib import Path
from typing import Any, Iterable, Sequence

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.marker import DataPoint
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError as error:  # pragma: no cover - environment gate
    raise SystemExit(
        "openpyxl is required. Install benchmark/requirements.txt first."
    ) from error


ALGORITHM_ORDER = [
    "tdtree_cpu_32t",
    "ours_cpu_32t",
    "ours_gopher_cpu_32t",
    "ours_gopher_gpu",
]
ALGORITHM_LABELS = {
    "tdtree_cpu_32t": "TD-Tree CPU 32T",
    "ours_cpu_32t": "Ours CPU 32T",
    "ours_gopher_cpu_32t": "Ours+Gopher CPU 32T",
    "ours_gopher_gpu": "Ours+Gopher GPU",
}
DATASET_ORDER = [
    "sx_superuser",
    "wiki_talk",
    "sx_stackoverflow",
    "bitcoin",
]
DATASET_LABELS = {
    "sx_superuser": "Superuser",
    "wiki_talk": "Wiki-Talk",
    "sx_stackoverflow": "Stack Overflow",
    "bitcoin": "Bitcoin",
}
ROOT = Path(__file__).resolve().parents[1]
BASELINE_ALGORITHM = "tdtree_cpu_32t"
PROPOSED_ALGORITHMS = tuple(
    algorithm for algorithm in ALGORITHM_ORDER if algorithm != BASELINE_ALGORITHM
)
STABLE_LABEL_MODE = "stable-id"
PALETTE = ["4472C4", "ED7D31", "A5A5A5", "70AD47"]
EXPECTED_SHEETS = [
    "Overview",
    "Dataset Comparison",
    "Dataset Pivot",
    "Match Results",
    "Raw Runs",
    "Configuration",
    "Methodology",
]
FILTER_SHEET_NAMES = [
    "Filter Materialization",
    "Filter Raw Runs",
]
EXPECTED_SHEETS_WITH_FILTER = [
    "Overview",
    "Dataset Comparison",
    "Dataset Pivot",
    "Filter Materialization",
    "Match Results",
    "Raw Runs",
    "Filter Raw Runs",
    "Configuration",
    "Methodology",
]
FILTER_RAW_COLUMNS = [
    "dataset",
    "repeat",
    "execution_position",
    "status",
    "exit_code",
    "wall_s",
    "filter_core_ms",
    "serialize_loop_ms",
    "unattributed_read_group_close_ms",
    "raw_unique_edges",
    "retained_unique_edges",
    "generated_rows",
    "generated_bytes",
    "generated_sha256",
    "canonical_validation",
    "canonical_generated_sha256",
    "canonical_supplied_sha256",
    "max_rss_kb",
    "major_faults",
    "minor_faults",
    "voluntary_context_switches",
    "involuntary_context_switches",
    "cpu_affinity",
    "raw_path",
    "supplied_filtered_path",
    "source_sha256",
    "executable_sha256",
    "started_at_utc",
    "ended_at_utc",
]

TITLE_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
SUBHEADER_FILL = PatternFill("solid", fgColor="E2F0D9")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
PASS_FILL = PatternFill("solid", fgColor="E2F0D9")
FAIL_FILL = PatternFill("solid", fgColor="F4CCCC")
THIN_GREY = Side(style="thin", color="B7B7B7")


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create the focused dataset-wise 32T/GPU Excel workbook."
    )
    parser.add_argument("--input-dir", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--validate-only", type=Path)
    parser.add_argument(
        "--filter-materialization-dir",
        type=Path,
        help=(
            "Optional directory containing the raw-to-_filtered runs.csv, "
            "summary.csv, validation.json, and environment.json."
        ),
    )
    return parser.parse_args(argv)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def numeric(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() in {"none", "null", "nan"}:
        return None
    try:
        result = float(text)
    except ValueError:
        return None
    return result if math.isfinite(result) else None


def integer(value: Any) -> int | None:
    number = numeric(value)
    if number is None or not number.is_integer():
        return None
    return int(number)


def truthy(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"true", "1", "yes", "on"}:
        return True
    if text in {"false", "0", "no", "off"}:
        return False
    return None


def first_numeric(row: dict[str, str], keys: Iterable[str]) -> float | None:
    for key in keys:
        value = numeric(row.get(key))
        if value is not None:
            return value
    return None


def percentile(values: Sequence[float], fraction: float) -> float:
    ordered = sorted(values)
    if len(ordered) == 1:
        return ordered[0]
    index = (len(ordered) - 1) * fraction
    lower = math.floor(index)
    upper = math.ceil(index)
    if lower == upper:
        return ordered[lower]
    weight = index - lower
    return ordered[lower] * (1.0 - weight) + ordered[upper] * weight


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def flatten_json(value: Any, prefix: str = "") -> list[tuple[str, Any]]:
    rows: list[tuple[str, Any]] = []
    if isinstance(value, dict):
        for key, child in value.items():
            child_prefix = f"{prefix}.{key}" if prefix else str(key)
            rows.extend(flatten_json(child, child_prefix))
    elif isinstance(value, list):
        rows.append((prefix, json.dumps(value, ensure_ascii=False)))
    else:
        rows.append((prefix, value))
    return rows


def case_algorithm_value(
    case: dict[str, Any],
    key: str,
    algorithm: str,
) -> Any:
    overrides = case.get(f"{key}_by_algorithm", {})
    if overrides is None:
        overrides = {}
    if not isinstance(overrides, dict):
        raise ValueError(f"{key}_by_algorithm must be an object")
    return overrides.get(algorithm, case.get(key))


def resolved_manifest_path(value: Any) -> Path:
    if value is None or not str(value).strip():
        raise ValueError("Manifest data path is blank")
    expanded = str(value).replace("{root}", str(ROOT))
    return Path(expanded).resolve()


def option_value(flags: Sequence[Any], option: str) -> str | None:
    values = [str(value) for value in flags]
    result: str | None = None
    for index, flag in enumerate(values):
        if flag == option:
            if index + 1 >= len(values):
                raise ValueError(f"{option} requires a value")
            result = values[index + 1]
        elif flag.startswith(f"{option}="):
            result = flag.split("=", 1)[1]
    return result


def validate_input_protocol(
    manifest: dict[str, Any],
    case: dict[str, Any],
) -> None:
    algorithms = manifest.get("algorithms", {})
    data_overrides = case.get("data_by_algorithm")
    variant_overrides = case.get("input_variant_by_algorithm")
    if not isinstance(data_overrides, dict):
        raise ValueError(
            f"{case.get('id')}: data_by_algorithm must explicitly map "
            f"{BASELINE_ALGORITHM} to the raw input"
        )
    if BASELINE_ALGORITHM not in data_overrides:
        raise ValueError(
            f"{case.get('id')}: data_by_algorithm is missing "
            f"{BASELINE_ALGORITHM}"
        )
    if not isinstance(variant_overrides, dict):
        raise ValueError(
            f"{case.get('id')}: input_variant_by_algorithm must explicitly "
            f"describe the raw baseline"
        )
    if BASELINE_ALGORITHM not in variant_overrides:
        raise ValueError(
            f"{case.get('id')}: input_variant_by_algorithm is missing "
            f"{BASELINE_ALGORITHM}"
        )

    baseline_path = resolved_manifest_path(
        case_algorithm_value(case, "data", BASELINE_ALGORITHM)
    )
    proposed_paths = {
        algorithm: resolved_manifest_path(
            case_algorithm_value(case, "data", algorithm)
        )
        for algorithm in PROPOSED_ALGORITHMS
    }
    if "_filtered" in baseline_path.stem:
        raise ValueError(
            f"{case.get('id')}: Original (TD-Tree) must use the raw input"
        )
    if any("_filtered" not in path.stem for path in proposed_paths.values()):
        raise ValueError(
            f"{case.get('id')}: every proposed method must use a "
            "precomputed _filtered input"
        )
    if any(path == baseline_path for path in proposed_paths.values()):
        raise ValueError(
            f"{case.get('id')}: raw baseline and precomputed-filtered "
            "proposed inputs must differ"
        )
    if len(set(proposed_paths.values())) != 1:
        raise ValueError(
            f"{case.get('id')}: proposed methods must share one filtered input"
        )

    baseline_variant = str(
        case_algorithm_value(
            case, "input_variant", BASELINE_ALGORITHM
        )
        or ""
    ).lower()
    if "raw" not in baseline_variant or "filtered" in baseline_variant:
        raise ValueError(
            f"{case.get('id')}: baseline input_variant must identify raw input"
        )
    for algorithm in PROPOSED_ALGORITHMS:
        variant = str(
            case_algorithm_value(case, "input_variant", algorithm) or ""
        ).lower()
        if "filtered" not in variant:
            raise ValueError(
                f"{case.get('id')}/{algorithm}: input_variant must identify "
                "precomputed filtered input"
            )

    for algorithm in ALGORITHM_ORDER:
        label_mode = option_value(
            algorithms.get(algorithm, {}).get("flags", []),
            "--label-mode",
        )
        if label_mode != STABLE_LABEL_MODE:
            raise ValueError(
                f"{case.get('id')}/{algorithm}: manifest must request "
                f"--label-mode {STABLE_LABEL_MODE}"
            )


def style_title(ws: Any, title: str, subtitle: str | None = None) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)
    cell = ws.cell(1, 1, title)
    cell.fill = TITLE_FILL
    cell.font = Font(color="FFFFFF", bold=True, size=16)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=12)
        subtitle_cell = ws.cell(2, 1, subtitle)
        subtitle_cell.font = Font(color="666666", italic=True)
        subtitle_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 32
        return 3
    return 2


def style_header(ws: Any, row: int, start: int, end: int) -> None:
    for column in range(start, end + 1):
        cell = ws.cell(row, column)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True, color="1F1F1F")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GREY)


def add_table(ws: Any, name: str, min_row: int, max_row: int, max_col: int) -> None:
    if max_row <= min_row:
        return
    table = Table(
        displayName=name,
        ref=f"A{min_row}:{get_column_letter(max_col)}{max_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def autosize(ws: Any, maximum: int = 42) -> None:
    for column_cells in ws.columns:
        length = 0
        for cell in column_cells:
            if cell.value is not None:
                length = max(length, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
            maximum, max(10, length + 2)
        )


def gather(
    input_dir: Path,
) -> tuple[
    dict[tuple[str, str], dict[str, Any]],
    list[dict[str, str]],
    dict[str, Any],
    dict[str, Any],
]:
    runs_path = input_dir / "runs.csv"
    correctness_path = input_dir / "correctness.csv"
    manifest_path = input_dir / "manifest.used.json"
    for path in (runs_path, correctness_path, manifest_path):
        if not path.exists():
            raise ValueError(f"Missing required input: {path}")

    runs = read_csv(runs_path)
    correctness_rows = read_csv(correctness_path)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    environment_path = input_dir / "environment.json"
    environment = (
        json.loads(environment_path.read_text(encoding="utf-8"))
        if environment_path.exists()
        else {}
    )

    cases_by_id = {
        str(case["id"]): case for case in manifest.get("cases", [])
    }
    expected_cases = {
        case_id: str(case.get("dataset_id", ""))
        for case_id, case in cases_by_id.items()
    }
    if set(expected_cases.values()) != set(DATASET_ORDER):
        raise ValueError(
            "The manifest must contain exactly the four expected dataset IDs."
        )
    manifest_algorithms = list(manifest.get("algorithms", {}).keys())
    if manifest_algorithms != ALGORITHM_ORDER:
        raise ValueError(
            f"Algorithm order mismatch: expected {ALGORITHM_ORDER}, "
            f"found {manifest_algorithms}"
        )
    for case in cases_by_id.values():
        validate_input_protocol(manifest, case)

    correctness = {
        row["case_id"]: row.get("status", "")
        for row in correctness_rows
        if row.get("semantic_group") == "consecutive"
    }
    grouped: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for row in runs:
        case_id = row.get("case_id", "")
        algorithm = row.get("algorithm", "")
        if case_id in expected_cases and algorithm in ALGORITHM_ORDER:
            grouped[(case_id, algorithm)].append(row)

    summaries: dict[tuple[str, str], dict[str, Any]] = {}
    for case_id, dataset_id in expected_cases.items():
        case = cases_by_id[case_id]
        if correctness.get(case_id) != "pass":
            raise ValueError(f"Correctness did not pass for {case_id}")
        case_counts: set[int] = set()
        for algorithm in ALGORITHM_ORDER:
            group = grouped.get((case_id, algorithm), [])
            expected_repeats = int(
                case.get("repeats", manifest.get("repeats", 1))
            )
            if len(group) != expected_repeats:
                raise ValueError(
                    f"{case_id}/{algorithm}: expected {expected_repeats} runs, "
                    f"found {len(group)}"
                )
            if any(row.get("status") != "ok" for row in group):
                raise ValueError(f"{case_id}/{algorithm}: non-ok run")

            expected_data_path = resolved_manifest_path(
                case_algorithm_value(case, "data", algorithm)
            )
            recorded_data_values = {
                str(row.get("data", "")).strip() for row in group
            }
            if "" in recorded_data_values or len(recorded_data_values) != 1:
                raise ValueError(
                    f"{case_id}/{algorithm}: missing or unstable recorded "
                    "data path"
                )
            recorded_data_value = next(iter(recorded_data_values))
            recorded_data_path = Path(recorded_data_value).resolve()
            if recorded_data_path != expected_data_path:
                raise ValueError(
                    f"{case_id}/{algorithm}: recorded data path "
                    f"{recorded_data_path} does not match manifest "
                    f"{expected_data_path}"
                )

            expected_input_variant = str(
                case_algorithm_value(case, "input_variant", algorithm) or ""
            )
            recorded_input_variants = {
                str(row.get("input_variant", "")).strip() for row in group
            }
            if recorded_input_variants != {expected_input_variant}:
                raise ValueError(
                    f"{case_id}/{algorithm}: recorded input_variant "
                    f"{sorted(recorded_input_variants)} does not match "
                    f"manifest {expected_input_variant!r}"
                )

            recorded_data_hashes = {
                str(row.get("data_sha256", "")).strip() for row in group
            }
            if "" in recorded_data_hashes or len(recorded_data_hashes) != 1:
                raise ValueError(
                    f"{case_id}/{algorithm}: missing or unstable data_sha256"
                )
            recorded_data_hash = next(iter(recorded_data_hashes))

            recorded_label_modes = {
                str(row.get("label_mode_config", "")).strip() for row in group
            }
            if recorded_label_modes != {STABLE_LABEL_MODE}:
                raise ValueError(
                    f"{case_id}/{algorithm}: every run must record "
                    f"label_mode_config={STABLE_LABEL_MODE}"
                )

            walls = [numeric(row.get("wall_s")) for row in group]
            if any(value is None or value <= 0 for value in walls):
                raise ValueError(f"{case_id}/{algorithm}: invalid wall time")
            wall_values = [float(value) for value in walls if value is not None]

            counts = {integer(row.get("match_count")) for row in group}
            if None in counts or len(counts) != 1:
                raise ValueError(f"{case_id}/{algorithm}: unstable/missing count")
            match_count = next(iter(counts))
            assert match_count is not None
            case_counts.add(match_count)

            thread_values = [
                first_numeric(
                    row,
                    (
                        "threads_used",
                        "result_threads_used",
                        "gopher_threads_used",
                    ),
                )
                for row in group
            ]
            worker_values = [
                first_numeric(
                    row,
                    (
                        "workers_with_tasks",
                        "result_workers_with_tasks",
                        "parallel_workers_with_tasks",
                        "gopher_parallel_statistics_workers_with_tasks",
                    ),
                )
                for row in group
            ]
            fallbacks = [
                truthy(
                    row.get("fallback_used")
                    or row.get("result_fallback_used")
                    or row.get("gopher_fallback_used")
                )
                for row in group
            ]
            if algorithm != "ours_gopher_gpu":
                if any(value is None for value in thread_values):
                    raise ValueError(
                        f"{case_id}/{algorithm}: missing actual threads_used"
                    )
                if any(
                    int(float(value)) != 32
                    for value in thread_values
                    if value is not None
                ):
                    raise ValueError(
                        f"{case_id}/{algorithm}: actual threads_used is not 32"
                    )
            elif any(value is True for value in fallbacks):
                raise ValueError(f"{case_id}/{algorithm}: CPU fallback observed")

            input_occurrences = next(
                (
                    integer(row.get("input_occurrences"))
                    for row in group
                    if integer(row.get("input_occurrences")) is not None
                ),
                None,
            )
            active_vertices = next(
                (
                    integer(row.get("active_vertices"))
                    for row in group
                    if integer(row.get("active_vertices")) is not None
                ),
                None,
            )
            summaries[(dataset_id, algorithm)] = {
                "case_id": case_id,
                "dataset_id": dataset_id,
                "algorithm": algorithm,
                "input_role": (
                    "raw baseline"
                    if algorithm == BASELINE_ALGORITHM
                    else "precomputed filtered proposed"
                ),
                "input_variant": expected_input_variant,
                "manifest_data": str(
                    case_algorithm_value(case, "data", algorithm)
                ),
                "data_path": str(recorded_data_path),
                "data_sha256": recorded_data_hash,
                "label_mode": STABLE_LABEL_MODE,
                "runs": len(group),
                "match_count": match_count,
                "wall_median": statistics.median(wall_values),
                "wall_q1": percentile(wall_values, 0.25),
                "wall_q3": percentile(wall_values, 0.75),
                "wall_min": min(wall_values),
                "wall_max": max(wall_values),
                "threads_used_median": (
                    statistics.median(
                        float(value)
                        for value in thread_values
                        if value is not None
                    )
                    if any(value is not None for value in thread_values)
                    else None
                ),
                "workers_with_tasks_median": (
                    statistics.median(
                        float(value)
                        for value in worker_values
                        if value is not None
                    )
                    if any(value is not None for value in worker_values)
                    else None
                ),
                "input_occurrences": input_occurrences,
                "active_vertices": active_vertices,
                "correctness": correctness[case_id],
                "fallback": any(value is True for value in fallbacks),
            }
        if len(case_counts) != 1:
            raise ValueError(
                f"{case_id}: implementations disagree on match count: "
                f"{sorted(case_counts)}"
            )
        proposed_paths = {
            summaries[(dataset_id, algorithm)]["data_path"]
            for algorithm in PROPOSED_ALGORITHMS
        }
        proposed_hashes = {
            summaries[(dataset_id, algorithm)]["data_sha256"]
            for algorithm in PROPOSED_ALGORITHMS
        }
        if len(proposed_paths) != 1 or len(proposed_hashes) != 1:
            raise ValueError(
                f"{case_id}: proposed methods did not record one shared "
                "precomputed filtered input"
            )
        if (
            summaries[(dataset_id, BASELINE_ALGORITHM)]["data_sha256"]
            in proposed_hashes
        ):
            raise ValueError(
                f"{case_id}: raw and filtered inputs unexpectedly share "
                "one data SHA-256"
            )
    return summaries, runs, manifest, environment


def gather_filter_materialization(
    filter_dir: Path,
    online_summaries: dict[tuple[str, str], dict[str, Any]],
) -> tuple[
    dict[str, dict[str, Any]],
    list[dict[str, str]],
    dict[str, Any],
]:
    required_paths = {
        "runs": filter_dir / "runs.csv",
        "summary": filter_dir / "summary.csv",
        "validation": filter_dir / "validation.json",
        "environment": filter_dir / "environment.json",
    }
    for path in required_paths.values():
        if not path.exists():
            raise ValueError(f"Missing required filter input: {path}")

    filter_runs = read_csv(required_paths["runs"])
    filter_summary_rows = read_csv(required_paths["summary"])
    validation = json.loads(
        required_paths["validation"].read_text(encoding="utf-8")
    )
    filter_environment = json.loads(
        required_paths["environment"].read_text(encoding="utf-8")
    )
    if not filter_runs:
        raise ValueError("Filter materialization runs.csv is empty")
    missing_columns = set(FILTER_RAW_COLUMNS) - set(filter_runs[0])
    if missing_columns:
        raise ValueError(
            "Filter materialization runs.csv is missing columns: "
            f"{sorted(missing_columns)}"
        )

    online_repeat_counts = {
        int(item["runs"]) for item in online_summaries.values()
    }
    if len(online_repeat_counts) != 1:
        raise ValueError("Online summaries have inconsistent repetition counts")
    expected_repeats = next(iter(online_repeat_counts))
    recorded_repeats = integer(filter_environment.get("repetitions"))
    if recorded_repeats != expected_repeats:
        raise ValueError(
            "Filter materialization repetition count does not match the "
            f"online benchmark: {recorded_repeats} vs {expected_repeats}"
        )
    expected_run_count = len(DATASET_ORDER) * expected_repeats
    if len(filter_runs) != expected_run_count:
        raise ValueError(
            "Filter materialization run count mismatch: expected "
            f"{expected_run_count}, found {len(filter_runs)}"
        )

    run_datasets = {row.get("dataset", "") for row in filter_runs}
    if run_datasets != set(DATASET_ORDER):
        raise ValueError(
            "Filter materialization dataset mismatch: expected "
            f"{DATASET_ORDER}, found {sorted(run_datasets)}"
        )
    identities: set[tuple[str, int]] = set()
    for row in filter_runs:
        dataset = row.get("dataset", "")
        repeat = integer(row.get("repeat"))
        position = integer(row.get("execution_position"))
        if repeat is None or not 1 <= repeat <= expected_repeats:
            raise ValueError(f"{dataset}: invalid filter repeat {repeat}")
        identity = (dataset, repeat)
        if identity in identities:
            raise ValueError(f"Duplicate filter run identity: {identity}")
        identities.add(identity)
        if position is None or not 1 <= position <= len(DATASET_ORDER):
            raise ValueError(
                f"{dataset}/r{repeat}: invalid execution position {position}"
            )
        if row.get("status") != "ok" or integer(row.get("exit_code")) != 0:
            raise ValueError(f"{dataset}/r{repeat}: non-ok filter run")
        if row.get("canonical_validation") != "pass":
            raise ValueError(
                f"{dataset}/r{repeat}: canonical validation did not pass"
            )

        wall_s = numeric(row.get("wall_s"))
        filter_core_ms = numeric(row.get("filter_core_ms"))
        serialize_loop_ms = numeric(row.get("serialize_loop_ms"))
        remainder_ms = numeric(
            row.get("unattributed_read_group_close_ms")
        )
        max_rss_kb = numeric(row.get("max_rss_kb"))
        if (
            wall_s is None
            or wall_s <= 0
            or filter_core_ms is None
            or filter_core_ms < 0
            or serialize_loop_ms is None
            or serialize_loop_ms < 0
            or remainder_ms is None
            or remainder_ms < 0
            or max_rss_kb is None
            or max_rss_kb <= 0
        ):
            raise ValueError(
                f"{dataset}/r{repeat}: invalid filter timing or RSS value"
            )
        expected_remainder = max(
            0.0,
            wall_s * 1000.0 - filter_core_ms - serialize_loop_ms,
        )
        if not math.isclose(
            remainder_ms,
            expected_remainder,
            rel_tol=0.0,
            abs_tol=0.001,
        ):
            raise ValueError(
                f"{dataset}/r{repeat}: unattributed remainder is inconsistent "
                "with full wall minus internal timers"
            )
        for key in (
            "raw_unique_edges",
            "retained_unique_edges",
            "generated_rows",
            "generated_bytes",
        ):
            value = integer(row.get(key))
            if value is None or value < 0:
                raise ValueError(
                    f"{dataset}/r{repeat}: invalid integer field {key}"
                )
        if integer(row["retained_unique_edges"]) > integer(
            row["raw_unique_edges"]
        ):
            raise ValueError(
                f"{dataset}/r{repeat}: retained edges exceed raw edges"
            )
        for key in (
            "raw_path",
            "supplied_filtered_path",
            "source_sha256",
            "executable_sha256",
            "generated_sha256",
        ):
            if not str(row.get(key, "")).strip():
                raise ValueError(
                    f"{dataset}/r{repeat}: blank filter provenance field {key}"
                )

    expected_identities = {
        (dataset, repeat)
        for dataset in DATASET_ORDER
        for repeat in range(1, expected_repeats + 1)
    }
    if identities != expected_identities:
        raise ValueError("Filter materialization run identities are incomplete")
    for repeat in range(1, expected_repeats + 1):
        positions = {
            integer(row.get("execution_position"))
            for row in filter_runs
            if integer(row.get("repeat")) == repeat
        }
        if positions != set(range(1, len(DATASET_ORDER) + 1)):
            raise ValueError(
                f"Filter materialization r{repeat} execution positions "
                "are incomplete"
            )

    cpu_affinities = {
        str(row.get("cpu_affinity", "")).strip() for row in filter_runs
    }
    environment_affinity = str(
        filter_environment.get("cpu_affinity", "")
    ).strip()
    if (
        len(cpu_affinities) != 1
        or environment_affinity not in cpu_affinities
        or not environment_affinity.isdigit()
    ):
        raise ValueError(
            "Filter materialization must record one stable single-CPU affinity"
        )
    source_hashes = {
        str(row.get("source_sha256", "")).strip() for row in filter_runs
    }
    executable_hashes = {
        str(row.get("executable_sha256", "")).strip()
        for row in filter_runs
    }
    if source_hashes != {
        str(filter_environment.get("source_sha256", "")).strip()
    }:
        raise ValueError("Filter source SHA-256 is missing or unstable")
    if executable_hashes != {
        str(filter_environment.get("executable_sha256", "")).strip()
    }:
        raise ValueError("Filter executable SHA-256 is missing or unstable")
    timing_boundaries = filter_environment.get("timing_boundaries", {})
    if not isinstance(timing_boundaries, dict) or not {
        "wall_s",
        "filter_core_ms",
        "serialize_loop_ms",
    }.issubset(timing_boundaries):
        raise ValueError("Filter environment is missing timing boundaries")
    if not str(filter_environment.get("page_cache_policy", "")).strip():
        raise ValueError("Filter environment is missing page-cache policy")

    if validation.get("status") != "pass":
        raise ValueError("Filter materialization validation.json did not pass")
    if integer(validation.get("observed_runs")) != expected_run_count:
        raise ValueError("Filter validation has an incorrect observed run count")
    if integer(validation.get("expected_runs")) != expected_run_count:
        raise ValueError("Filter validation has an incorrect expected run count")
    if integer(validation.get("non_ok_runs")) != 0:
        raise ValueError("Filter validation reports non-ok runs")
    validation_datasets = validation.get("datasets", {})
    if (
        not isinstance(validation_datasets, dict)
        or set(validation_datasets) != set(DATASET_ORDER)
    ):
        raise ValueError("Filter validation dataset coverage is incomplete")

    summary_by_dataset = {
        row.get("dataset", ""): row for row in filter_summary_rows
    }
    if set(summary_by_dataset) != set(DATASET_ORDER):
        raise ValueError("Filter summary dataset coverage is incomplete")
    if len(filter_summary_rows) != len(DATASET_ORDER):
        raise ValueError("Filter summary contains duplicate dataset rows")

    summary_numeric_fields = (
        "wall_s",
        "filter_core_ms",
        "serialize_loop_ms",
        "unattributed_read_group_close_ms",
        "max_rss_kb",
    )
    result: dict[str, dict[str, Any]] = {}
    for dataset in DATASET_ORDER:
        group = [
            row for row in filter_runs if row.get("dataset") == dataset
        ]
        summary_row = summary_by_dataset[dataset]
        if integer(summary_row.get("repetitions")) != expected_repeats:
            raise ValueError(
                f"{dataset}: filter summary repetition count is incorrect"
            )
        if summary_row.get("canonical_validation") != "pass":
            raise ValueError(
                f"{dataset}: filter summary canonical validation failed"
            )

        stable_integer_fields = (
            "raw_unique_edges",
            "retained_unique_edges",
            "generated_rows",
            "generated_bytes",
        )
        stable_values: dict[str, int] = {}
        for field in stable_integer_fields:
            values = {integer(row.get(field)) for row in group}
            if None in values or len(values) != 1:
                raise ValueError(
                    f"{dataset}: unstable filter field {field}"
                )
            value = next(iter(values))
            assert value is not None
            stable_values[field] = value
            if integer(summary_row.get(field)) != value:
                raise ValueError(
                    f"{dataset}: filter summary disagrees on {field}"
                )

        recomputed: dict[str, float] = {}
        for field in summary_numeric_fields:
            values = [numeric(row.get(field)) for row in group]
            if any(value is None for value in values):
                raise ValueError(
                    f"{dataset}: missing numeric filter field {field}"
                )
            clean_values = [
                float(value) for value in values if value is not None
            ]
            expected_values = {
                "median": statistics.median(clean_values),
                "min": min(clean_values),
                "max": max(clean_values),
            }
            for statistic_name, expected_value in expected_values.items():
                recorded_value = numeric(
                    summary_row.get(f"{field}_{statistic_name}")
                )
                if recorded_value is None or not math.isclose(
                    recorded_value,
                    expected_value,
                    rel_tol=1e-12,
                    abs_tol=1e-9,
                ):
                    raise ValueError(
                        f"{dataset}: filter summary disagrees on "
                        f"{field}_{statistic_name}"
                    )
                recomputed[f"{field}_{statistic_name}"] = expected_value

        raw_paths = {
            str(Path(row["raw_path"]).resolve()) for row in group
        }
        supplied_paths = {
            str(Path(row["supplied_filtered_path"]).resolve())
            for row in group
        }
        generated_hashes = {
            str(row.get("generated_sha256", "")).strip() for row in group
        }
        if len(raw_paths) != 1 or len(supplied_paths) != 1:
            raise ValueError(f"{dataset}: unstable filter input paths")
        if len(generated_hashes) != 1:
            raise ValueError(f"{dataset}: unstable generated byte SHA-256")
        raw_path = next(iter(raw_paths))
        supplied_path = next(iter(supplied_paths))
        baseline = online_summaries[(dataset, BASELINE_ALGORITHM)]
        proposed_items = [
            online_summaries[(dataset, algorithm)]
            for algorithm in PROPOSED_ALGORITHMS
        ]
        if (
            baseline["input_occurrences"] is None
            or baseline["input_occurrences"] <= 0
        ):
            raise ValueError(
                f"{dataset}: baseline input occurrence count is unavailable"
            )
        if raw_path != str(Path(baseline["data_path"]).resolve()):
            raise ValueError(
                f"{dataset}: filter raw path does not match TD-Tree input"
            )
        if any(
            supplied_path != str(Path(item["data_path"]).resolve())
            for item in proposed_items
        ):
            raise ValueError(
                f"{dataset}: supplied filtered path does not match proposed "
                "online input"
            )
        proposed_occurrences = {
            item["input_occurrences"]
            for item in proposed_items
            if item["input_occurrences"] is not None
        }
        if proposed_occurrences != {stable_values["generated_rows"]}:
            raise ValueError(
                f"{dataset}: generated rows do not match proposed online input"
            )

        validation_item = validation_datasets[dataset]
        if validation_item.get("status") != "pass":
            raise ValueError(
                f"{dataset}: canonical content validation did not pass"
            )
        canonical_generated_hash = str(
            validation_item.get("generated_sha256", "")
        ).strip()
        canonical_supplied_hash = str(
            validation_item.get("supplied_sha256", "")
        ).strip()
        if (
            not canonical_generated_hash
            or canonical_generated_hash != canonical_supplied_hash
        ):
            raise ValueError(
                f"{dataset}: canonical generated and supplied hashes differ"
            )
        observed_canonical_pairs = {
            (
                str(row.get("canonical_generated_sha256", "")).strip(),
                str(row.get("canonical_supplied_sha256", "")).strip(),
            )
            for row in group
            if str(row.get("canonical_generated_sha256", "")).strip()
            or str(row.get("canonical_supplied_sha256", "")).strip()
        }
        if observed_canonical_pairs != {
            (canonical_generated_hash, canonical_supplied_hash)
        }:
            raise ValueError(
                f"{dataset}: canonical hashes are missing or inconsistent "
                "between runs.csv and validation.json"
            )
        if integer(validation_item.get("generated_rows")) != stable_values[
            "generated_rows"
        ]:
            raise ValueError(
                f"{dataset}: validation generated-row count is inconsistent"
            )
        if integer(validation_item.get("generated_bytes")) != stable_values[
            "generated_bytes"
        ]:
            raise ValueError(
                f"{dataset}: validation generated-byte count is inconsistent"
            )

        result[dataset] = {
            "dataset": dataset,
            "runs": expected_repeats,
            "cpu_affinity": environment_affinity,
            **stable_values,
            **recomputed,
            "raw_path": raw_path,
            "supplied_filtered_path": supplied_path,
            "raw_data_sha256": baseline["data_sha256"],
            "supplied_filtered_data_sha256": proposed_items[0][
                "data_sha256"
            ],
            "generated_sha256": next(iter(generated_hashes)),
            "canonical_content_sha256": canonical_generated_hash,
            "canonical_validation": "pass",
            "source_sha256": next(iter(source_hashes)),
            "executable_sha256": next(iter(executable_hashes)),
        }
    return result, filter_runs, filter_environment


def build_workbook(
    input_dir: Path,
    output: Path,
    filter_materialization_dir: Path | None = None,
) -> dict[str, Any]:
    summaries, runs, manifest, environment = gather(input_dir)
    filter_summaries: dict[str, dict[str, Any]] | None = None
    filter_runs: list[dict[str, str]] = []
    filter_environment: dict[str, Any] = {}
    if filter_materialization_dir is not None:
        (
            filter_summaries,
            filter_runs,
            filter_environment,
        ) = gather_filter_materialization(
            filter_materialization_dir,
            summaries,
        )
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Dataset-wise TD-Tree 32T Baseline Comparison"
    workbook.properties.subject = (
        "TD-Tree CPU 32T vs Ours CPU 32T vs Ours+Gopher CPU 32T vs GPU"
    )

    # Overview
    ws = workbook.create_sheet("Overview")
    row = style_title(
        ws,
        "Dataset-wise 32T / GPU Performance Evaluation",
        "Protocol: Original (TD-Tree) reads raw data; proposed methods read "
        "precomputed filtered data; stable-ID labels, path3, k=3.",
    )
    overview_rows = [
        (
            "Validation",
            "PASS",
            "All runs completed, counts agree, and recorded per-algorithm "
            "inputs match the manifest protocol.",
        ),
        ("Datasets", 4, "Superuser, Wiki-Talk, Stack Overflow, Bitcoin"),
        ("Algorithms", 4, "TD-Tree 32T baseline plus three requested methods"),
        ("Runs", len(runs), "Five isolated repetitions per dataset/method"),
        (
            "Zero-match datasets",
            sum(
                1
                for dataset in DATASET_ORDER
                if summaries[(dataset, ALGORITHM_ORDER[0])]["match_count"] == 0
            ),
            "Count-only suppresses mapping materialization, not match counting.",
        ),
        ("Baseline", "TD-Tree CPU 32T", "Explicit consecutive compatibility mode"),
        (
            "Input protocol",
            "TD-Tree=raw; proposed=precomputed filtered",
            "Each run records its resolved path, input variant, and SHA-256.",
        ),
        (
            "Label mode",
            STABLE_LABEL_MODE,
            "External-ID-derived labels stay invariant across raw and filtered "
            "vertex universes.",
        ),
        (
            "Primary metric",
            "Median online wall time (s)",
            "Includes loading and online preprocessing of the selected input; "
            "excludes construction of the proposed filtered files.",
        ),
        (
            "Timing boundary",
            "Precomputed-filter online comparison",
            "Not a numerical replication of the 260408_cdh paper total-time "
            "figures, which include filtering construction overhead.",
        ),
        ("GPU", "Resident continuation", "No silent CPU fallback"),
    ]
    if filter_summaries is not None:
        overview_rows.extend(
            [
                (
                    "Filter validation",
                    "PASS",
                    f"{len(filter_runs)}/{len(filter_runs)} single-CPU "
                    "materialization runs completed and all four canonical "
                    "generated/supplied outputs agree.",
                ),
                (
                    "Filter materialization runs",
                    len(filter_runs),
                    "Five rotated process runs per dataset on one logical CPU.",
                ),
                (
                    "Filter timing primary",
                    "Full raw-to-filtered wall time",
                    "Includes raw read/parse/group, internal filter work, "
                    "serialization, close, and process overhead.",
                ),
                (
                    "Filter internal timers",
                    "Core and serialize loop",
                    "Narrow diagnostic timers; neither is the complete "
                    "materialization cost.",
                ),
                (
                    "Filter page cache",
                    filter_environment.get("page_cache_policy", ""),
                    "The benchmark rotates dataset order but does not flush "
                    "the operating-system page cache.",
                ),
            ]
        )
    ws.append(["Metric", "Value", "Interpretation"])
    style_header(ws, row, 1, 3)
    for item in overview_rows:
        ws.append(item)
    for data_row in range(row + 1, row + 1 + len(overview_rows)):
        if ws.cell(data_row, 1).value in {
            "Validation",
            "Filter validation",
        }:
            ws.cell(data_row, 2).fill = PASS_FILL
            ws.cell(data_row, 2).font = Font(bold=True, color="006100")
        if ws.cell(data_row, 1).value == "Zero-match datasets":
            ws.cell(data_row, 2).fill = PASS_FILL
            ws.cell(data_row, 2).font = Font(bold=True, color="006100")
    ws.freeze_panes = f"A{row + 1}"
    autosize(ws)

    # Long comparison table
    ws = workbook.create_sheet("Dataset Comparison")
    header_row = style_title(
        ws,
        "Dataset-by-dataset performance",
        "Speedup = TD-Tree CPU 32T median / method median. "
        "The denominator reads raw data; proposed methods read precomputed "
        "filtered data. Values above 1 are faster in this online protocol.",
    )
    headers = [
        "Dataset",
        "Case ID",
        "Algorithm",
        "Backend",
        "Requested CPU threads",
        "Used CPU threads (median)",
        "Workers with tasks (median)",
        "Match count",
        "Match status",
        "Median wall (s)",
        "Q1 wall (s)",
        "Q3 wall (s)",
        "IQR wall (s)",
        "Min wall (s)",
        "Max wall (s)",
        "Speedup vs TD-Tree 32T",
        "Runs",
        "Correctness",
        "CPU fallback",
    ]
    ws.append(headers)
    style_header(ws, header_row, 1, len(headers))
    for dataset in DATASET_ORDER:
        baseline = summaries[(dataset, ALGORITHM_ORDER[0])]["wall_median"]
        for algorithm in ALGORITHM_ORDER:
            item = summaries[(dataset, algorithm)]
            backend = "GPU" if algorithm == "ours_gopher_gpu" else "CPU"
            requested = 1 if backend == "GPU" else 32
            ws.append(
                [
                    DATASET_LABELS[dataset],
                    item["case_id"],
                    ALGORITHM_LABELS[algorithm],
                    backend,
                    requested,
                    item["threads_used_median"],
                    item["workers_with_tasks_median"],
                    item["match_count"],
                    "NONZERO" if item["match_count"] > 0 else "ZERO",
                    item["wall_median"],
                    item["wall_q1"],
                    item["wall_q3"],
                    item["wall_q3"] - item["wall_q1"],
                    item["wall_min"],
                    item["wall_max"],
                    baseline / item["wall_median"],
                    item["runs"],
                    item["correctness"].upper(),
                    item["fallback"],
                ]
            )
    last_row = ws.max_row
    for data_row in range(header_row + 1, last_row + 1):
        for column in range(10, 17):
            ws.cell(data_row, column).number_format = "0.0000"
        ws.cell(data_row, 8).number_format = "#,##0"
        ws.cell(data_row, 9).fill = (
            PASS_FILL if ws.cell(data_row, 9).value == "NONZERO" else WARNING_FILL
        )
        ws.cell(data_row, 18).fill = PASS_FILL
    add_table(ws, "DatasetComparisonTable", header_row, last_row, len(headers))
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
    autosize(ws)

    # Pivot and runtime charts
    ws = workbook.create_sheet("Dataset Pivot")
    header_row = style_title(
        ws,
        "Runtime and speedup summary",
        "Exact values are cold-process online medians in seconds; proposed "
        "filter-construction time is outside this workbook's boundary.",
    )
    headers = [
        "Dataset",
        "TD-Tree raw input rows",
        "Match count",
        "TD-Tree CPU 32T (s)",
        "Ours CPU 32T (s)",
        "Ours+Gopher CPU 32T (s)",
        "Ours+Gopher GPU (s)",
        "TD-Tree baseline speedup",
        "Ours speedup",
        "Ours+Gopher 32T speedup",
        "GPU speedup",
    ]
    ws.append(headers)
    style_header(ws, header_row, 1, len(headers))
    for dataset in DATASET_ORDER:
        baseline = summaries[(dataset, ALGORITHM_ORDER[0])]["wall_median"]
        runtimes = [
            summaries[(dataset, algorithm)]["wall_median"]
            for algorithm in ALGORITHM_ORDER
        ]
        ws.append(
            [
                DATASET_LABELS[dataset],
                summaries[(dataset, ALGORITHM_ORDER[0])]["input_occurrences"],
                summaries[(dataset, ALGORITHM_ORDER[0])]["match_count"],
                *runtimes,
                1.0,
                baseline / runtimes[1],
                baseline / runtimes[2],
                baseline / runtimes[3],
            ]
        )
    for data_row in range(header_row + 1, ws.max_row + 1):
        ws.cell(data_row, 2).number_format = "#,##0"
        ws.cell(data_row, 3).number_format = "#,##0"
        for column in range(4, 12):
            ws.cell(data_row, column).number_format = "0.0000"

    for offset, dataset in enumerate(DATASET_ORDER):
        source_row = header_row + 1 + offset
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = f"{DATASET_LABELS[dataset]} end-to-end runtime"
        chart.y_axis.title = "Seconds (median)"
        chart.x_axis.title = "Method"
        chart.height = 7.0
        chart.width = 12.0
        data = Reference(
            ws, min_col=4, max_col=7, min_row=source_row, max_row=source_row
        )
        categories = Reference(
            ws, min_col=4, max_col=7, min_row=header_row, max_row=header_row
        )
        chart.add_data(data, from_rows=True, titles_from_data=False)
        chart.set_categories(categories)
        chart.legend = None
        chart.dLbls = DataLabelList()
        chart.dLbls.showVal = True
        chart.dLbls.numFmt = "0.000"
        if chart.series:
            for index, color in enumerate(PALETTE):
                point = DataPoint(idx=index)
                point.graphicalProperties.solidFill = color
                point.graphicalProperties.line.solidFill = color
                chart.series[0].dPt.append(point)
        anchor_column = "A" if offset % 2 == 0 else "J"
        anchor_row = 10 + (offset // 2) * 16
        ws.add_chart(chart, f"{anchor_column}{anchor_row}")

    speedup_chart = BarChart()
    speedup_chart.type = "col"
    speedup_chart.style = 10
    speedup_chart.title = "Speedup vs TD-Tree CPU 32T"
    speedup_chart.y_axis.title = "Speedup (×)"
    speedup_chart.x_axis.title = "Dataset"
    speedup_chart.height = 8.0
    speedup_chart.width = 16.0
    speedup_data = Reference(
        ws,
        min_col=9,
        max_col=11,
        min_row=header_row,
        max_row=header_row + len(DATASET_ORDER),
    )
    speedup_categories = Reference(
        ws,
        min_col=1,
        min_row=header_row + 1,
        max_row=header_row + len(DATASET_ORDER),
    )
    speedup_chart.add_data(speedup_data, titles_from_data=True)
    speedup_chart.set_categories(speedup_categories)
    speedup_chart.dLbls = DataLabelList()
    speedup_chart.dLbls.showVal = False
    for series, color in zip(speedup_chart.series, PALETTE[1:]):
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.solidFill = color
    ws.add_chart(speedup_chart, "A42")
    ws.freeze_panes = f"A{header_row + 1}"
    autosize(ws)

    # Optional raw-to-_filtered materialization evidence.
    if filter_summaries is not None:
        ws = workbook.create_sheet("Filter Materialization")
        header_row = style_title(
            ws,
            "Raw-to-_filtered materialization time",
            "Primary metric is full process wall time on one logical CPU. "
            "Filter-core and serialization-loop values are narrower internal "
            "diagnostics; the unattributed remainder is not a direct read-time "
            "measurement.",
        )
        headers = [
            "Dataset",
            "Runs",
            "CPU affinity",
            "Full wall median (s)",
            "Full wall min (s)",
            "Full wall max (s)",
            "Filter-core median (ms)",
            "Filter-core min (ms)",
            "Filter-core max (ms)",
            "Serialize-loop median (ms)",
            "Serialize-loop min (ms)",
            "Serialize-loop max (ms)",
            "Unattributed remainder median (ms)",
            "Unattributed remainder min (ms)",
            "Unattributed remainder max (ms)",
            "Raw unique edge pairs",
            "Retained unique edge pairs",
            "Pair retention",
            "Raw occurrence rows",
            "Generated filtered rows",
            "Row retention",
            "Generated output (MiB)",
            "Peak RSS median (GiB)",
            "Peak RSS max (GiB)",
            "Canonical content match",
        ]
        ws.append(headers)
        style_header(ws, header_row, 1, len(headers))
        for dataset in DATASET_ORDER:
            item = filter_summaries[dataset]
            raw_occurrences = summaries[
                (dataset, BASELINE_ALGORITHM)
            ]["input_occurrences"]
            ws.append(
                [
                    DATASET_LABELS[dataset],
                    item["runs"],
                    item["cpu_affinity"],
                    item["wall_s_median"],
                    item["wall_s_min"],
                    item["wall_s_max"],
                    item["filter_core_ms_median"],
                    item["filter_core_ms_min"],
                    item["filter_core_ms_max"],
                    item["serialize_loop_ms_median"],
                    item["serialize_loop_ms_min"],
                    item["serialize_loop_ms_max"],
                    item[
                        "unattributed_read_group_close_ms_median"
                    ],
                    item["unattributed_read_group_close_ms_min"],
                    item["unattributed_read_group_close_ms_max"],
                    item["raw_unique_edges"],
                    item["retained_unique_edges"],
                    (
                        item["retained_unique_edges"]
                        / item["raw_unique_edges"]
                        if item["raw_unique_edges"]
                        else None
                    ),
                    raw_occurrences,
                    item["generated_rows"],
                    (
                        item["generated_rows"] / raw_occurrences
                        if raw_occurrences
                        else None
                    ),
                    item["generated_bytes"] / (1024.0**2),
                    item["max_rss_kb_median"] / (1024.0**2),
                    item["max_rss_kb_max"] / (1024.0**2),
                    item["canonical_validation"].upper(),
                ]
            )
        summary_last_row = ws.max_row
        for data_row in range(header_row + 1, summary_last_row + 1):
            for column in range(4, 16):
                ws.cell(data_row, column).number_format = "0.000"
            for column in (16, 17, 19, 20):
                ws.cell(data_row, column).number_format = "#,##0"
            for column in (18, 21):
                ws.cell(data_row, column).number_format = "0.0%"
            for column in range(22, 25):
                ws.cell(data_row, column).number_format = "0.000"
            ws.cell(data_row, 25).fill = PASS_FILL
            ws.cell(data_row, 25).font = Font(bold=True, color="006100")
        add_table(
            ws,
            "FilterMaterializationSummaryTable",
            header_row,
            summary_last_row,
            len(headers),
        )

        section_row = ws.max_row + 3
        ws.merge_cells(
            start_row=section_row,
            start_column=1,
            end_row=section_row,
            end_column=10,
        )
        section_cell = ws.cell(
            section_row,
            1,
            "One-shot and reuse-amortized interpretation",
        )
        section_cell.fill = SUBHEADER_FILL
        section_cell.font = Font(bold=True)
        section_cell.alignment = Alignment(vertical="center")
        amortization_header = section_row + 1
        amortization_headers = [
            "Dataset",
            "Method",
            "TD-Tree raw online median (s)",
            "Filter wall median (s)",
            "Proposed filtered online median (s)",
            "One-shot total (s)",
            "One-shot speedup vs TD-Tree",
            "10-query amortized per-query (s)",
            "10-query amortized speedup",
            "Break-even query count",
        ]
        for column, value in enumerate(amortization_headers, start=1):
            ws.cell(amortization_header, column, value)
        style_header(
            ws,
            amortization_header,
            1,
            len(amortization_headers),
        )
        for dataset in DATASET_ORDER:
            materialization = filter_summaries[dataset]["wall_s_median"]
            baseline = summaries[
                (dataset, BASELINE_ALGORITHM)
            ]["wall_median"]
            for algorithm in PROPOSED_ALGORITHMS:
                online = summaries[(dataset, algorithm)]["wall_median"]
                one_shot = materialization + online
                amortized_ten = online + materialization / 10.0
                break_even = (
                    math.ceil(materialization / (baseline - online))
                    if baseline > online
                    else None
                )
                ws.append(
                    [
                        DATASET_LABELS[dataset],
                        ALGORITHM_LABELS[algorithm],
                        baseline,
                        materialization,
                        online,
                        one_shot,
                        baseline / one_shot,
                        amortized_ten,
                        baseline / amortized_ten,
                        break_even,
                    ]
                )
        amortization_last_row = ws.max_row
        for data_row in range(
            amortization_header + 1,
            amortization_last_row + 1,
        ):
            for column in range(3, 10):
                ws.cell(data_row, column).number_format = "0.0000"
            ws.cell(data_row, 10).number_format = "0"
        add_table(
            ws,
            "FilterAmortizationTable",
            amortization_header,
            amortization_last_row,
            len(amortization_headers),
        )
        note_row = ws.max_row + 2
        ws.merge_cells(
            start_row=note_row,
            start_column=1,
            end_row=note_row,
            end_column=10,
        )
        note_cell = ws.cell(
            note_row,
            1,
            "Modeled reuse only: one-shot = M + O; 10-query amortized "
            "per-query = O + M/10; break-even is the smallest integer N with "
            "M + N·O ≤ N·B. M is materialization wall, O is proposed online "
            "median, and B is TD-Tree online median. This is not a measured "
            "ten-query batch.",
        )
        note_cell.font = Font(color="666666", italic=True)
        note_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[note_row].height = 42
        ws.freeze_panes = f"A{header_row + 1}"
        autosize(ws, maximum=38)

    # Match counts
    ws = workbook.create_sheet("Match Results")
    header_row = style_title(
        ws,
        "Matching-result audit",
        "All four methods produced the same nonzero count for each dataset. "
        "Count-only avoids writing every mapping but preserves the exact count.",
    )
    headers = [
        "Dataset",
        "TD-Tree CPU 32T",
        "Ours CPU 32T",
        "Ours+Gopher CPU 32T",
        "Ours+Gopher GPU",
        "All methods agree",
        "Result status",
    ]
    ws.append(headers)
    style_header(ws, header_row, 1, len(headers))
    for dataset in DATASET_ORDER:
        counts = [
            summaries[(dataset, algorithm)]["match_count"]
            for algorithm in ALGORITHM_ORDER
        ]
        ws.append(
            [
                DATASET_LABELS[dataset],
                *counts,
                len(set(counts)) == 1,
                "NONZERO" if counts[0] > 0 else "ZERO",
            ]
        )
    for data_row in range(header_row + 1, ws.max_row + 1):
        for column in range(2, 6):
            ws.cell(data_row, column).number_format = "#,##0"
        ws.cell(data_row, 6).fill = PASS_FILL
        ws.cell(data_row, 7).fill = (
            PASS_FILL if ws.cell(data_row, 7).value == "NONZERO" else WARNING_FILL
        )
    add_table(ws, "MatchResultsTable", header_row, ws.max_row, len(headers))

    match_chart = BarChart()
    match_chart.type = "bar"
    match_chart.style = 10
    match_chart.title = "Match count by dataset (log scale)"
    match_chart.x_axis.title = "Matches"
    match_chart.y_axis.title = "Dataset"
    match_chart.height = 8.0
    match_chart.width = 14.0
    match_chart.x_axis.scaling.logBase = 10
    match_chart.x_axis.scaling.min = 1
    match_data = Reference(
        ws,
        min_col=2,
        max_col=2,
        min_row=header_row,
        max_row=header_row + len(DATASET_ORDER),
    )
    match_categories = Reference(
        ws,
        min_col=1,
        min_row=header_row + 1,
        max_row=header_row + len(DATASET_ORDER),
    )
    match_chart.add_data(match_data, titles_from_data=True)
    match_chart.set_categories(match_categories)
    match_chart.legend = None
    match_chart.dLbls = DataLabelList()
    match_chart.dLbls.showVal = True
    match_chart.dLbls.numFmt = "#,##0"
    if match_chart.series:
        match_chart.series[0].graphicalProperties.solidFill = PALETTE[0]
        match_chart.series[0].graphicalProperties.line.solidFill = PALETTE[0]
    ws.add_chart(match_chart, "A11")
    ws.freeze_panes = f"A{header_row + 1}"
    autosize(ws)

    # Raw runs
    ws = workbook.create_sheet("Raw Runs")
    header_row = style_title(
        ws,
        "Raw benchmark repetitions",
        "One row per isolated process, including the resolved raw/filtered "
        "input path, variant, SHA-256, and stable label mode.",
    )
    raw_columns = [
        "case_id",
        "dataset_id",
        "algorithm",
        "input_variant",
        "data",
        "data_sha256",
        "label_mode_config",
        "repeat",
        "status",
        "wall_s",
        "match_count",
        "threads_requested",
        "threads_used",
        "workers_with_tasks",
        "parallel_workers_with_tasks",
        "cpu_affinity",
        "fallback_used",
        "gpu_variant",
        "gpuKernel_ms",
        "gpuReadyToResult_ms",
        "resource_max_rss_kb",
        "executable_sha256",
        "query_sha256",
        "run_signature",
    ]
    ws.append(raw_columns)
    style_header(ws, header_row, 1, len(raw_columns))
    selected_runs = [
        row
        for row in runs
        if row.get("algorithm") in ALGORITHM_ORDER
        and row.get("dataset_id") in DATASET_ORDER
    ]
    selected_runs.sort(
        key=lambda row: (
            DATASET_ORDER.index(row["dataset_id"]),
            ALGORITHM_ORDER.index(row["algorithm"]),
            int(row["repeat"]),
        )
    )
    for run in selected_runs:
        ws.append([run.get(column, "") for column in raw_columns])
    add_table(ws, "RawRunsTable32T", header_row, ws.max_row, len(raw_columns))
    ws.freeze_panes = f"A{header_row + 1}"
    autosize(ws, maximum=34)

    if filter_summaries is not None:
        ws = workbook.create_sheet("Filter Raw Runs")
        header_row = style_title(
            ws,
            "Raw filter-materialization repetitions",
            "One row per single-CPU process. generated_sha256 is the direct "
            "byte hash of the temporary output; canonical hashes validate "
            "content independent of row order.",
        )
        ws.append(FILTER_RAW_COLUMNS)
        style_header(ws, header_row, 1, len(FILTER_RAW_COLUMNS))
        sorted_filter_runs = sorted(
            filter_runs,
            key=lambda row: (
                DATASET_ORDER.index(row["dataset"]),
                int(row["repeat"]),
            ),
        )
        for run in sorted_filter_runs:
            ws.append([run.get(column, "") for column in FILTER_RAW_COLUMNS])
        add_table(
            ws,
            "FilterRawRunsTable",
            header_row,
            ws.max_row,
            len(FILTER_RAW_COLUMNS),
        )
        ws.freeze_panes = f"A{header_row + 1}"
        autosize(ws, maximum=34)

    # Configuration
    ws = workbook.create_sheet("Configuration")
    header_row = style_title(
        ws,
        "Exact configurations and provenance",
        "Original (TD-Tree) uses raw inputs; proposed methods use paired "
        "precomputed filtered inputs. All runs use stable-ID labels.",
    )
    headers = [
        "Algorithm ID",
        "Display label",
        "Backend",
        "Requested threads",
        "Role",
        "Semantics",
        "Label mode",
        "Executable",
        "Executable SHA-256",
        "Flags",
        "Interpretation boundary",
    ]
    ws.append(headers)
    style_header(ws, header_row, 1, len(headers))
    manifest_algorithms = manifest["algorithms"]
    representative_run = {
        algorithm: next(
            row for row in runs if row.get("algorithm") == algorithm
        )
        for algorithm in ALGORITHM_ORDER
    }
    boundaries = {
        "tdtree_cpu_32t": (
            "Engineering 32T root-task extension; consecutive mode, not "
            "paper-default cardinality; reads the raw graph."
        ),
        "ours_cpu_32t": (
            "Engineering 32T root-task extension of 260408_cdh; reads the "
            "precomputed filtered graph."
        ),
        "ours_gopher_cpu_32t": (
            "Candidate-DAG, reuse order, bounded memo, pruning, dynamic prefix; "
            "reads the precomputed filtered graph."
        ),
        "ours_gopher_gpu": (
            "Resident GPU Candidate-DAG; precomputed filtered graph, no CPU "
            "runtime memo, and no fallback."
        ),
    }
    for algorithm in ALGORITHM_ORDER:
        config = manifest_algorithms[algorithm]
        run = representative_run[algorithm]
        ws.append(
            [
                algorithm,
                ALGORITHM_LABELS[algorithm],
                "GPU" if algorithm == "ours_gopher_gpu" else "CPU",
                config.get("threads"),
                config.get("role"),
                config.get("semantic_group"),
                run.get("label_mode_config", ""),
                run.get("command", "").split()[0] if run.get("command") else "",
                run.get("executable_sha256", ""),
                " ".join(str(flag) for flag in config.get("flags", [])),
                boundaries[algorithm],
            ]
        )
    add_table(
        ws,
        "AlgorithmConfigurationTable32T",
        header_row,
        ws.max_row,
        len(headers),
    )
    dataset_header = ws.max_row + 3
    ws.cell(dataset_header, 1, "Dataset ID")
    ws.cell(dataset_header, 2, "Display")
    ws.cell(dataset_header, 3, "Algorithm")
    ws.cell(dataset_header, 4, "Input role")
    ws.cell(dataset_header, 5, "Input variant")
    ws.cell(dataset_header, 6, "Manifest data path")
    ws.cell(dataset_header, 7, "Recorded resolved data path")
    ws.cell(dataset_header, 8, "Recorded data SHA-256")
    ws.cell(dataset_header, 9, "Label mode")
    ws.cell(dataset_header, 10, "Query path")
    ws.cell(dataset_header, 11, "k")
    ws.cell(dataset_header, 12, "Observed match count")
    ws.cell(dataset_header, 13, "Path matches manifest")
    style_header(ws, dataset_header, 1, 13)
    for case in manifest["cases"]:
        dataset = case["dataset_id"]
        for algorithm in ALGORITHM_ORDER:
            item = summaries[(dataset, algorithm)]
            ws.append(
                [
                    dataset,
                    DATASET_LABELS[dataset],
                    algorithm,
                    item["input_role"],
                    item["input_variant"],
                    item["manifest_data"],
                    item["data_path"],
                    item["data_sha256"],
                    item["label_mode"],
                    case["query"],
                    case["k"],
                    item["match_count"],
                    str(resolved_manifest_path(item["manifest_data"]))
                    == item["data_path"],
                ]
            )
    add_table(
        ws,
        "DatasetInputProtocolTable32T",
        dataset_header,
        ws.max_row,
        13,
    )
    ws.freeze_panes = f"A{header_row + 1}"
    autosize(ws, maximum=52)

    # Methodology/environment
    ws = workbook.create_sheet("Methodology")
    header_row = style_title(
        ws,
        "Methodology, environment, and interpretation limits",
        "The workbook keeps performance, correctness, and configuration "
        "evidence on separate auditable sheets.",
    )
    ws.append(["Category", "Key", "Value"])
    style_header(ws, header_row, 1, 3)
    for key, value in manifest.get("methodology", {}).items():
        ws.append(["Methodology", key, value])
    for key, value in flatten_json(environment):
        ws.append(["Environment", key, value])
    if filter_summaries is not None:
        for key, value in flatten_json(filter_environment):
            ws.append(["Filter environment", key, value])
    limits = [
        (
            "Protocol",
            "per_algorithm_input",
            "Original (TD-Tree) reads each unfiltered raw file; Ours, "
            "Ours+Gopher, and Ours+Gopher-GPU read the paired precomputed "
            "filtered file.",
        ),
        (
            "Protocol",
            "stable_id_labels",
            "Every run uses --label-mode stable-id so a vertex keeps the same "
            "synthetic label when raw and filtered vertex universes differ.",
        ),
        (
            "Boundary",
            "precomputed_filter_online",
            "Reported wall time includes loading and online preprocessing of "
            "the selected input, but excludes construction of the proposed "
            "filtered files.",
        ),
        (
            "Boundary",
            "paper_total_time",
            "The 260408_cdh paper's filtering-overhead-inclusive total-time "
            "results have a different boundary; this workbook is not their "
            "numerical replication.",
        ),
        (
            "Limit",
            "count_only",
            "Individual mapping lines are not materialized; exact match counts are retained.",
        ),
        (
            "Limit",
            "cold_process",
            "GPU results include WSL/ROCDXG process and context startup.",
        ),
        (
            "Limit",
            "gpu_memo",
            "The GPU executor does not implement the CPU dependency-key runtime memo tables.",
        ),
        (
            "Limit",
            "query_population",
            "Queries are fixed concrete-witness queries, not the unpublished paper query population.",
        ),
    ]
    if filter_summaries is not None:
        limits.extend(
            [
                (
                    "Filter boundary",
                    "filter_materialization_wall",
                    "Primary offline filtering time is full raw-to-_filtered "
                    "process wall time: launch, raw read/parse/group, filter "
                    "core, serialization, close, and process overhead.",
                ),
                (
                    "Filter boundary",
                    "filter_core_internal",
                    "Internal timer starts after raw read/group and covers "
                    "per-edge timestamp sort/dedup, consecutive-pair "
                    "selection, and filtered-map insertion.",
                ),
                (
                    "Filter boundary",
                    "filter_serialize_internal",
                    "Internal output formatting/write-loop timer; output close "
                    "is outside this timer.",
                ),
                (
                    "Filter boundary",
                    "filter_unattributed_remainder",
                    "Full wall minus the two internal timers. It combines "
                    "raw read/parse/group, output close, process overhead, and "
                    "unmeasured work; it is not direct read time.",
                ),
                (
                    "Filter protocol",
                    "filter_single_cpu",
                    "The repository materializer is pinned to one logical CPU; "
                    f"recorded affinity is "
                    f"{filter_environment.get('cpu_affinity', '')}.",
                ),
                (
                    "Filter protocol",
                    "filter_page_cache",
                    str(
                        filter_environment.get("page_cache_policy", "")
                    ),
                ),
                (
                    "Filter validation",
                    "canonical_content_validation",
                    "Generated and supplied _filtered files are numerically "
                    "sorted before hashing; all four canonical hashes agree. "
                    "Direct byte hashes may differ because row order differs.",
                ),
                (
                    "Filter lineage",
                    "filter_input_path_linkage",
                    "Materialization raw/supplied paths exactly match the "
                    "TD-Tree/proposed online paths. Online run SHA-256 values "
                    "identify those files; the materialization CSV itself "
                    "does not record their direct byte hashes.",
                ),
                (
                    "Modeled interpretation",
                    "filter_one_shot_formula",
                    "One-shot proposed total = M + O.",
                ),
                (
                    "Modeled interpretation",
                    "filter_amortized_formula",
                    "N-query amortized proposed time per query = O + M/N. "
                    "The displayed N=10 values are modeled, not a measured "
                    "ten-query batch.",
                ),
                (
                    "Modeled interpretation",
                    "filter_break_even_formula",
                    "When B > O, break-even query count is "
                    "ceil(M / (B - O)); otherwise no finite break-even exists.",
                ),
            ]
        )
    for limit in limits:
        ws.append(limit)
    add_table(ws, "MethodologyTable32T", header_row, ws.max_row, 3)
    ws.freeze_panes = f"A{header_row + 1}"
    autosize(ws, maximum=80)

    for ws in workbook.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        for row_cells in ws.iter_rows():
            for cell in row_cells:
                if cell.value is not None:
                    alignment = copy(cell.alignment)
                    alignment.vertical = alignment.vertical or "top"
                    if alignment.wrap_text is None:
                        alignment.wrap_text = False
                    cell.alignment = alignment

    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=f".{output.stem}.", suffix=".xlsx", dir=output.parent, delete=False
    ) as temporary_handle:
        temporary = Path(temporary_handle.name)
    try:
        workbook.save(temporary)
        validate_workbook(
            temporary,
            expected_filter_materialization=filter_summaries is not None,
        )
        os.replace(temporary, output)
    finally:
        if temporary.exists():
            temporary.unlink()
    return validate_workbook(
        output,
        expected_filter_materialization=filter_summaries is not None,
    )


def validate_workbook(
    path: Path,
    expected_filter_materialization: bool | None = None,
) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False)
    observed_filter_sheets = {
        sheet for sheet in FILTER_SHEET_NAMES if sheet in workbook.sheetnames
    }
    if expected_filter_materialization is None:
        expected_filter_materialization = bool(observed_filter_sheets)
    expected_sheets = (
        EXPECTED_SHEETS_WITH_FILTER
        if expected_filter_materialization
        else EXPECTED_SHEETS
    )
    if workbook.sheetnames != expected_sheets:
        raise ValueError(
            f"Sheet mismatch: expected {expected_sheets}, "
            f"found {workbook.sheetnames}"
        )
    comparison = workbook["Dataset Comparison"]
    header_row = next(
        row
        for row in range(1, comparison.max_row + 1)
        if comparison.cell(row, 1).value == "Dataset"
    )
    rows = [
        row
        for row in comparison.iter_rows(
            min_row=header_row + 1, max_row=comparison.max_row, values_only=True
        )
        if row[0]
    ]
    if len(rows) != 16:
        raise ValueError(f"Expected 16 comparison rows, found {len(rows)}")
    if any(row[7] is None for row in rows):
        raise ValueError("A comparison row has a blank match count")
    if any(row[17] != "PASS" for row in rows):
        raise ValueError("A comparison row did not pass correctness")
    match_ws = workbook["Match Results"]
    match_header = next(
        row
        for row in range(1, match_ws.max_row + 1)
        if match_ws.cell(row, 1).value == "Dataset"
    )
    match_rows = [
        row
        for row in match_ws.iter_rows(
            min_row=match_header + 1, max_row=match_ws.max_row, values_only=True
        )
        if row[0]
    ]
    if len(match_rows) != 4:
        raise ValueError(f"Expected four match-result rows, found {len(match_rows)}")
    if any(row[6] != "NONZERO" for row in match_rows):
        raise ValueError("A selected dataset has zero or missing matches")

    raw_ws = workbook["Raw Runs"]
    raw_header = next(
        row
        for row in range(1, raw_ws.max_row + 1)
        if raw_ws.cell(row, 1).value == "case_id"
    )
    raw_headers = {
        str(raw_ws.cell(raw_header, column).value): column
        for column in range(1, raw_ws.max_column + 1)
    }
    required_raw_columns = {
        "algorithm",
        "input_variant",
        "data",
        "data_sha256",
        "label_mode_config",
    }
    if not required_raw_columns.issubset(raw_headers):
        raise ValueError("Raw Runs is missing input-protocol provenance columns")
    raw_protocol_rows = [
        row
        for row in raw_ws.iter_rows(
            min_row=raw_header + 1,
            max_row=raw_ws.max_row,
            values_only=True,
        )
        if row[0]
    ]
    for row in raw_protocol_rows:
        algorithm = row[raw_headers["algorithm"] - 1]
        input_variant = str(row[raw_headers["input_variant"] - 1] or "")
        data_path = str(row[raw_headers["data"] - 1] or "")
        data_hash = str(row[raw_headers["data_sha256"] - 1] or "")
        label_mode = row[raw_headers["label_mode_config"] - 1]
        if not data_path or not data_hash:
            raise ValueError("Raw Runs has blank data path or SHA-256")
        if label_mode != STABLE_LABEL_MODE:
            raise ValueError("Raw Runs has a non-stable-ID label mode")
        if algorithm == BASELINE_ALGORITHM:
            if "_filtered" in Path(data_path).stem or "raw" not in input_variant:
                raise ValueError("Raw Runs baseline input is not raw")
        elif algorithm in PROPOSED_ALGORITHMS:
            if (
                "_filtered" not in Path(data_path).stem
                or "filtered" not in input_variant
            ):
                raise ValueError(
                    "Raw Runs proposed-method input is not precomputed filtered"
                )

    configuration = workbook["Configuration"]
    protocol_header = next(
        row
        for row in range(1, configuration.max_row + 1)
        if configuration.cell(row, 1).value == "Dataset ID"
    )
    protocol_rows = [
        row
        for row in configuration.iter_rows(
            min_row=protocol_header + 1,
            max_row=configuration.max_row,
            values_only=True,
        )
        if row[0]
    ]
    if len(protocol_rows) != len(DATASET_ORDER) * len(ALGORITHM_ORDER):
        raise ValueError(
            "Configuration does not contain all 16 per-algorithm input rows"
        )
    if any(row[8] != STABLE_LABEL_MODE for row in protocol_rows):
        raise ValueError("Configuration has a non-stable-ID label mode")
    if any(row[12] is not True for row in protocol_rows):
        raise ValueError(
            "Configuration contains a data path that does not match manifest"
        )

    overview = workbook["Overview"]
    overview_values = {
        str(overview.cell(row, 1).value): overview.cell(row, 2).value
        for row in range(1, overview.max_row + 1)
    }
    if overview_values.get("Label mode") != STABLE_LABEL_MODE:
        raise ValueError("Overview does not declare stable-ID labeling")
    if overview_values.get("Timing boundary") != (
        "Precomputed-filter online comparison"
    ):
        raise ValueError("Overview does not declare the online timing boundary")

    methodology = workbook["Methodology"]
    methodology_keys = {
        str(methodology.cell(row, 2).value)
        for row in range(1, methodology.max_row + 1)
    }
    required_methodology_keys = {
        "per_algorithm_input",
        "stable_id_labels",
        "precomputed_filter_online",
        "paper_total_time",
    }
    if not required_methodology_keys.issubset(methodology_keys):
        raise ValueError(
            "Methodology is missing an input or timing-boundary declaration"
        )

    filter_summary_row_count = 0
    filter_raw_run_count = 0
    filter_amortization_row_count = 0
    if expected_filter_materialization:
        filter_ws = workbook["Filter Materialization"]
        filter_header = next(
            row
            for row in range(1, filter_ws.max_row + 1)
            if filter_ws.cell(row, 1).value == "Dataset"
            and filter_ws.cell(row, 2).value == "Runs"
        )
        filter_headers = {
            str(filter_ws.cell(filter_header, column).value): column
            for column in range(1, filter_ws.max_column + 1)
            if filter_ws.cell(filter_header, column).value is not None
        }
        required_filter_headers = {
            "Dataset",
            "Runs",
            "CPU affinity",
            "Full wall median (s)",
            "Full wall min (s)",
            "Full wall max (s)",
            "Filter-core median (ms)",
            "Serialize-loop median (ms)",
            "Unattributed remainder median (ms)",
            "Raw unique edge pairs",
            "Retained unique edge pairs",
            "Raw occurrence rows",
            "Generated filtered rows",
            "Peak RSS median (GiB)",
            "Canonical content match",
        }
        if not required_filter_headers.issubset(filter_headers):
            raise ValueError(
                "Filter Materialization is missing required summary columns"
            )
        filter_summary_rows = [
            tuple(
                filter_ws.cell(row, column).value
                for column in range(1, filter_ws.max_column + 1)
            )
            for row in range(
                filter_header + 1,
                filter_header + 1 + len(DATASET_ORDER),
            )
        ]
        filter_summary_row_count = len(filter_summary_rows)
        if {
            row[filter_headers["Dataset"] - 1]
            for row in filter_summary_rows
        } != set(DATASET_LABELS.values()):
            raise ValueError(
                "Filter Materialization does not contain all four datasets"
            )
        expected_online_repeats = {
            row[16] for row in rows
        }
        if len(expected_online_repeats) != 1:
            raise ValueError("Comparison rows have unstable run counts")
        expected_filter_repeats = next(iter(expected_online_repeats))
        for row in filter_summary_rows:
            dataset_label = row[filter_headers["Dataset"] - 1]
            repetitions = integer(row[filter_headers["Runs"] - 1])
            wall_median = numeric(
                row[filter_headers["Full wall median (s)"] - 1]
            )
            wall_min = numeric(
                row[filter_headers["Full wall min (s)"] - 1]
            )
            wall_max = numeric(
                row[filter_headers["Full wall max (s)"] - 1]
            )
            filter_core = numeric(
                row[filter_headers["Filter-core median (ms)"] - 1]
            )
            serialize_loop = numeric(
                row[filter_headers["Serialize-loop median (ms)"] - 1]
            )
            remainder = numeric(
                row[
                    filter_headers[
                        "Unattributed remainder median (ms)"
                    ]
                    - 1
                ]
            )
            raw_edges = integer(
                row[filter_headers["Raw unique edge pairs"] - 1]
            )
            retained_edges = integer(
                row[filter_headers["Retained unique edge pairs"] - 1]
            )
            raw_rows = integer(
                row[filter_headers["Raw occurrence rows"] - 1]
            )
            generated_rows = integer(
                row[filter_headers["Generated filtered rows"] - 1]
            )
            peak_rss = numeric(
                row[filter_headers["Peak RSS median (GiB)"] - 1]
            )
            if (
                repetitions != expected_filter_repeats
                or wall_median is None
                or wall_min is None
                or wall_max is None
                or not 0 < wall_min <= wall_median <= wall_max
                or filter_core is None
                or filter_core < 0
                or serialize_loop is None
                or serialize_loop < 0
                or remainder is None
                or remainder < 0
                or raw_edges is None
                or retained_edges is None
                or not 0 <= retained_edges <= raw_edges
                or raw_rows is None
                or generated_rows is None
                or not 0 <= generated_rows <= raw_rows
                or peak_rss is None
                or peak_rss <= 0
            ):
                raise ValueError(
                    f"Filter Materialization has invalid values for "
                    f"{dataset_label}"
                )
            if row[
                filter_headers["Canonical content match"] - 1
            ] != "PASS":
                raise ValueError(
                    f"Filter Materialization canonical validation failed for "
                    f"{dataset_label}"
                )

        amortization_header = next(
            row
            for row in range(filter_header + 1, filter_ws.max_row + 1)
            if filter_ws.cell(row, 1).value == "Dataset"
            and filter_ws.cell(row, 2).value == "Method"
        )
        amortization_headers = {
            str(filter_ws.cell(amortization_header, column).value): column
            for column in range(1, filter_ws.max_column + 1)
            if filter_ws.cell(amortization_header, column).value is not None
        }
        required_amortization_headers = {
            "Dataset",
            "Method",
            "TD-Tree raw online median (s)",
            "Filter wall median (s)",
            "Proposed filtered online median (s)",
            "One-shot total (s)",
            "One-shot speedup vs TD-Tree",
            "10-query amortized per-query (s)",
            "10-query amortized speedup",
            "Break-even query count",
        }
        if not required_amortization_headers.issubset(
            amortization_headers
        ):
            raise ValueError(
                "Filter Materialization is missing amortization columns"
            )
        filter_amortization_rows = [
            tuple(
                filter_ws.cell(row, column).value
                for column in range(1, filter_ws.max_column + 1)
            )
            for row in range(
                amortization_header + 1,
                amortization_header
                + 1
                + len(DATASET_ORDER) * len(PROPOSED_ALGORITHMS),
            )
        ]
        filter_amortization_row_count = len(filter_amortization_rows)
        expected_pairs = {
            (DATASET_LABELS[dataset], ALGORITHM_LABELS[algorithm])
            for dataset in DATASET_ORDER
            for algorithm in PROPOSED_ALGORITHMS
        }
        observed_pairs = {
            (
                row[amortization_headers["Dataset"] - 1],
                row[amortization_headers["Method"] - 1],
            )
            for row in filter_amortization_rows
        }
        if observed_pairs != expected_pairs:
            raise ValueError(
                "Filter amortization table has incomplete dataset/method rows"
            )
        for row in filter_amortization_rows:
            baseline = numeric(
                row[
                    amortization_headers[
                        "TD-Tree raw online median (s)"
                    ]
                    - 1
                ]
            )
            materialization = numeric(
                row[
                    amortization_headers["Filter wall median (s)"]
                    - 1
                ]
            )
            online = numeric(
                row[
                    amortization_headers[
                        "Proposed filtered online median (s)"
                    ]
                    - 1
                ]
            )
            one_shot = numeric(
                row[amortization_headers["One-shot total (s)"] - 1]
            )
            one_shot_speedup = numeric(
                row[
                    amortization_headers[
                        "One-shot speedup vs TD-Tree"
                    ]
                    - 1
                ]
            )
            amortized_ten = numeric(
                row[
                    amortization_headers[
                        "10-query amortized per-query (s)"
                    ]
                    - 1
                ]
            )
            amortized_speedup = numeric(
                row[
                    amortization_headers[
                        "10-query amortized speedup"
                    ]
                    - 1
                ]
            )
            break_even = integer(
                row[
                    amortization_headers["Break-even query count"]
                    - 1
                ]
            )
            if (
                baseline is None
                or materialization is None
                or online is None
                or one_shot is None
                or one_shot_speedup is None
                or amortized_ten is None
                or amortized_speedup is None
            ):
                raise ValueError(
                    "Filter amortization table contains a blank numeric value"
                )
            expected_one_shot = materialization + online
            expected_amortized_ten = online + materialization / 10.0
            expected_break_even = (
                math.ceil(materialization / (baseline - online))
                if baseline > online
                else None
            )
            if (
                not math.isclose(
                    one_shot,
                    expected_one_shot,
                    rel_tol=1e-12,
                    abs_tol=1e-12,
                )
                or not math.isclose(
                    one_shot_speedup,
                    baseline / expected_one_shot,
                    rel_tol=1e-12,
                    abs_tol=1e-12,
                )
                or not math.isclose(
                    amortized_ten,
                    expected_amortized_ten,
                    rel_tol=1e-12,
                    abs_tol=1e-12,
                )
                or not math.isclose(
                    amortized_speedup,
                    baseline / expected_amortized_ten,
                    rel_tol=1e-12,
                    abs_tol=1e-12,
                )
                or break_even != expected_break_even
            ):
                raise ValueError(
                    "Filter one-shot/amortized calculations are inconsistent"
                )

        filter_raw_ws = workbook["Filter Raw Runs"]
        filter_raw_header = next(
            row
            for row in range(1, filter_raw_ws.max_row + 1)
            if filter_raw_ws.cell(row, 1).value == "dataset"
            and filter_raw_ws.cell(row, 2).value == "repeat"
        )
        filter_raw_headers = {
            str(filter_raw_ws.cell(filter_raw_header, column).value): column
            for column in range(1, filter_raw_ws.max_column + 1)
        }
        if not set(FILTER_RAW_COLUMNS).issubset(filter_raw_headers):
            raise ValueError("Filter Raw Runs is missing provenance columns")
        filter_raw_rows = [
            row
            for row in filter_raw_ws.iter_rows(
                min_row=filter_raw_header + 1,
                max_row=filter_raw_ws.max_row,
                values_only=True,
            )
            if row[0]
        ]
        filter_raw_run_count = len(filter_raw_rows)
        if filter_raw_run_count != (
            len(DATASET_ORDER) * int(expected_filter_repeats)
        ):
            raise ValueError(
                "Filter Raw Runs does not contain all expected repetitions"
            )
        filter_raw_identities: set[tuple[str, int]] = set()
        protocol_paths = {
            (row[0], row[2]): str(Path(row[6]).resolve())
            for row in protocol_rows
        }
        for row in filter_raw_rows:
            dataset = str(row[filter_raw_headers["dataset"] - 1])
            repeat = integer(row[filter_raw_headers["repeat"] - 1])
            if dataset not in DATASET_ORDER or repeat is None:
                raise ValueError("Filter Raw Runs has an invalid identity")
            identity = (dataset, repeat)
            if identity in filter_raw_identities:
                raise ValueError("Filter Raw Runs has a duplicate identity")
            filter_raw_identities.add(identity)
            if (
                row[filter_raw_headers["status"] - 1] != "ok"
                or integer(row[filter_raw_headers["exit_code"] - 1]) != 0
                or row[
                    filter_raw_headers["canonical_validation"] - 1
                ]
                != "pass"
            ):
                raise ValueError("Filter Raw Runs contains a failed run")
            affinity = str(
                row[filter_raw_headers["cpu_affinity"] - 1] or ""
            )
            if not affinity.isdigit():
                raise ValueError(
                    "Filter Raw Runs has a non-single-CPU affinity"
                )
            raw_path = str(
                Path(
                    str(row[filter_raw_headers["raw_path"] - 1])
                ).resolve()
            )
            supplied_path = str(
                Path(
                    str(
                        row[
                            filter_raw_headers[
                                "supplied_filtered_path"
                            ]
                            - 1
                        ]
                    )
                ).resolve()
            )
            if raw_path != protocol_paths[
                (dataset, BASELINE_ALGORITHM)
            ]:
                raise ValueError(
                    "Filter Raw Runs raw path does not match Configuration"
                )
            if any(
                supplied_path != protocol_paths[(dataset, algorithm)]
                for algorithm in PROPOSED_ALGORITHMS
            ):
                raise ValueError(
                    "Filter Raw Runs supplied path does not match Configuration"
                )

        if overview_values.get("Filter validation") != "PASS":
            raise ValueError("Overview does not declare filter validation")
        if overview_values.get("Filter timing primary") != (
            "Full raw-to-filtered wall time"
        ):
            raise ValueError(
                "Overview does not declare full materialization wall time"
            )
        required_filter_methodology_keys = {
            "filter_materialization_wall",
            "filter_core_internal",
            "filter_serialize_internal",
            "filter_unattributed_remainder",
            "canonical_content_validation",
            "filter_one_shot_formula",
            "filter_amortized_formula",
            "filter_break_even_formula",
        }
        if not required_filter_methodology_keys.issubset(
            methodology_keys
        ):
            raise ValueError(
                "Methodology is missing filter boundary or reuse definitions"
            )

    charts = sum(len(sheet._charts) for sheet in workbook.worksheets)
    if charts != 6:
        raise ValueError(f"Expected six charts, found {charts}")
    return {
        "validation": "passed",
        "path": str(path.resolve()),
        "sheets": len(workbook.sheetnames),
        "charts": charts,
        "comparison_rows": len(rows),
        "datasets": len(match_rows),
        "input_protocol_rows": len(protocol_rows),
        "zero_match_datasets": 0,
        "filter_materialization_rows": filter_summary_row_count,
        "filter_amortization_rows": filter_amortization_row_count,
        "filter_raw_runs": filter_raw_run_count,
    }


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        if args.validate_only:
            receipt = validate_workbook(args.validate_only.resolve())
        else:
            if args.input_dir is None or args.output is None:
                raise ValueError(
                    "--input-dir and --output are required unless --validate-only is used"
                )
            receipt = build_workbook(
                args.input_dir.resolve(),
                args.output.resolve(),
                (
                    args.filter_materialization_dir.resolve()
                    if args.filter_materialization_dir is not None
                    else None
                ),
            )
            receipt["sha256"] = sha256(args.output.resolve())
        print(json.dumps(receipt, ensure_ascii=False, indent=2, sort_keys=True))
        return 0
    except (OSError, ValueError, json.JSONDecodeError) as error:
        print(f"Dataset comparison workbook failed: {error}", file=os.sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
