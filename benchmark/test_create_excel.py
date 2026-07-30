#!/usr/bin/env python3
"""End-to-end tests for benchmark/create_excel.py using synthetic fixtures."""

from __future__ import annotations

import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook

import create_excel


HERE = Path(__file__).resolve().parent
SCRIPT = HERE / "create_excel.py"
FULL_FIXTURE = HERE / "fixtures" / "excel"
SPARSE_FIXTURE = HERE / "fixtures" / "excel_sparse"


class ExcelReportEndToEndTest(unittest.TestCase):
    maxDiff = None

    def test_thread_inference_accepts_runner_gopher_fields(self) -> None:
        self.assertEqual(
            create_excel.infer_threads(
                {"gopher_threads_requested": 16},
                "ours_gopher_reuse",
                "scaling_case",
            ),
            16,
        )
        self.assertEqual(
            create_excel.infer_threads(
                {"result_threads_requested": 32},
                "ours_gopher_reuse",
                "scaling_case",
            ),
            32,
        )
        self.assertEqual(
            create_excel.infer_threads({}, "original_consecutive", "case"),
            1,
        )
        self.assertEqual(create_excel.infer_threads({}, "ours", "case"), 1)

    def test_gpu_backend_and_timing_rows_are_inferred_without_mixing_units(
        self,
    ) -> None:
        self.assertEqual(
            create_excel.infer_backend({}, "ours_gopher_gpu"), "gpu"
        )
        self.assertEqual(create_excel.infer_backend({}, "ours"), "cpu")
        summaries = [
            {
                "case_id": "gpu_smoke",
                "semantic_group": "consecutive",
                "dataset": "synthetic",
                "query": "path3",
                "query_edges": 3,
                "k": 3,
                "algorithm": "original_consecutive",
                "algorithm_role": "baseline",
                "suite": "gpu_smoke",
                "threads": 1,
                "backend": "cpu",
                "device": "",
                "runtime_s": 2.0,
                "match_count": 7,
                "successful_runs": 3,
                "repeat_match_count_consistent": True,
                "source": {},
            },
            {
                "case_id": "gpu_smoke",
                "semantic_group": "consecutive",
                "dataset": "synthetic",
                "query": "path3",
                "query_edges": 3,
                "k": 3,
                "algorithm": "ours_gopher_none_paper_t1",
                "algorithm_role": "gpu_closest_cpu_reference_no_runtime_memo",
                "suite": "gpu_smoke",
                "threads": 1,
                "backend": "cpu",
                "device": "",
                "runtime_s": 1.25,
                "match_count": 7,
                "successful_runs": 3,
                "repeat_match_count_consistent": True,
                "source": {},
            },
            {
                "case_id": "gpu_smoke",
                "semantic_group": "consecutive",
                "dataset": "synthetic",
                "query": "path3",
                "query_edges": 3,
                "k": 3,
                "algorithm": "ours_gopher_reuse_t1",
                "algorithm_role": "cpu_single_thread",
                "suite": "gpu_smoke",
                "threads": 1,
                "backend": "cpu",
                "device": "",
                "runtime_s": 1.0,
                "match_count": 7,
                "successful_runs": 3,
                "repeat_match_count_consistent": True,
                "source": {},
            },
            {
                "case_id": "gpu_smoke",
                "semantic_group": "consecutive",
                "dataset": "synthetic",
                "query": "path3",
                "query_edges": 3,
                "k": 3,
                "algorithm": "ours_gopher_gpu",
                "algorithm_role": "gpu",
                "suite": "gpu_smoke",
                "threads": None,
                "backend": "rocm",
                "device": "gfx1151",
                "runtime_s": 0.5,
                "match_count": 7,
                "successful_runs": 3,
                "repeat_match_count_consistent": True,
                "source": {
                    "gopher_gpu_kernel_ms_median": 125.0,
                    "gopher_gpu_h2d_ms_median": 25.0,
                    "gopher_gpu_d2h_ms_median": 5.0,
                    "gopher_gpu_batches_median": 4,
                    "gopher_fallback_used_value": False,
                    "gopher_verification_requested_value": False,
                    "gopher_verification_passed_value": False,
                },
            },
        ]
        rows = create_excel.build_gpu_rows(
            summaries,
            [
                {
                    "case_id": "gpu_smoke",
                    "semantic_group": "consecutive",
                    "status": "pass",
                }
            ],
        )
        self.assertEqual(len(rows), 1)
        self.assertAlmostEqual(rows[0]["GPU kernel (s)"], 0.125)
        self.assertAlmostEqual(rows[0]["Host-to-device (s)"], 0.025)
        self.assertEqual(rows[0]["CPU 1T reference"], "ours_gopher_reuse_t1")
        self.assertEqual(
            rows[0]["Closest CPU no-memo reference"],
            "ours_gopher_none_paper_t1",
        )
        self.assertAlmostEqual(
            rows[0]["Speedup vs closest CPU no-memo (×)"], 2.5
        )
        self.assertAlmostEqual(rows[0]["Speedup vs CPU 1T (×)"], 2.0)
        self.assertTrue(rows[0]["Execution eligible"])

    def test_gpu_logical_reference_requires_matching_run_configuration(
        self,
    ) -> None:
        target = {
            "algorithm": "gpu",
            "backend": "rocm",
            "threads": None,
            "runtime_s": 0.5,
            "source": {
                "order_config_value": "reuse",
                "selectivity_config_value": "paper",
                "statistics_config_value": "off",
            },
        }
        entries = [
            target,
            {
                "algorithm": "wrong_fast_cpu",
                "algorithm_role": "no_runtime_memo",
                "backend": "cpu",
                "threads": 1,
                "runtime_s": 0.1,
                "source": {
                    "memo_config_value": "none",
                    "order_config_value": "reuse",
                    "selectivity_config_value": "intended",
                    "statistics_config_value": "off",
                },
            },
            {
                "algorithm": "exact_cpu",
                "algorithm_role": "no_runtime_memo",
                "backend": "cpu",
                "threads": 1,
                "runtime_s": 0.2,
                "source": {
                    "memo_config_value": "none",
                    "order_config_value": "reuse",
                    "selectivity_config_value": "paper",
                    "statistics_config_value": "off",
                },
            },
        ]
        selected = create_excel.choose_gpu_logical_cpu_reference(
            entries, target_gpu=target
        )
        self.assertIsNotNone(selected)
        self.assertEqual(selected["algorithm"], "exact_cpu")

    def test_numeric_candidate_metrics_are_not_misclassified_as_ids(self) -> None:
        self.assertEqual(
            create_excel.parse_csv_value(
                "generated_candidates_median", "19647168.0"
            ),
            19647168.0,
        )
        self.assertEqual(
            create_excel.parse_csv_value(
                "gopher_plan_id", "18220454072368105075"
            ),
            "18220454072368105075",
        )

    def test_missing_exact_metric_does_not_fall_back_to_unrelated_number(
        self,
    ) -> None:
        source = {
            "enumerate_matches_ns_median": 4_510_000,
            "parallel_worker_utilization_median": 0.9,
        }
        self.assertIsNone(
            create_excel.first_numeric_metric(
                source, ("gopher_gpu_kernel_launches_median",)
            )
        )

    def test_gpu_statistics_off_does_not_render_zero_work_as_measured(self) -> None:
        summaries = [
            {
                "case_id": "gpu_stats_off",
                "semantic_group": "consecutive",
                "dataset": "synthetic",
                "query": "path3",
                "query_edges": 3,
                "k": 3,
                "algorithm": "ours_gopher_gpu",
                "algorithm_role": "gpu",
                "suite": "gpu",
                "threads": None,
                "backend": "rocm",
                "device": "gfx1151",
                "runtime_s": 0.5,
                "match_count": 7,
                "successful_runs": 3,
                "repeat_match_count_consistent": True,
                "source": {
                    "gopher_statistics_enabled_value": False,
                    "gopher_statistics_generated_candidates_median": 0,
                    "gopher_statistics_accepted_candidates_median": 0,
                    "gopher_gpu_kernel_launches_median": 16,
                    "gopher_fallback_used_value": False,
                },
            }
        ]
        rows = create_excel.build_gpu_rows(
            summaries,
            [
                {
                    "case_id": "gpu_stats_off",
                    "semantic_group": "consecutive",
                    "status": "pass",
                }
            ],
        )
        self.assertEqual(len(rows), 1)
        self.assertIsNone(rows[0]["GPU work items"])
        self.assertIsNone(rows[0]["Accepted GPU candidates"])
        self.assertEqual(rows[0]["GPU batches"], 16)

    def test_gpu_depth_summary_preserves_unavailable_workload_fields(self) -> None:
        rows = create_excel.build_gpu_depth_summary_rows(
            [
                {
                    "case_id": "gpu_case",
                    "algorithm": "resident",
                    "gpu_variant": "resident_continuation",
                    "depth": 2,
                    "status": "ok",
                    "workload_metrics_available": False,
                    "input_states": 64,
                    "output_states": 128,
                    "generated_candidates": None,
                    "maximum_to_mean_load_ratio": None,
                    "grid_blocks": 320,
                }
            ]
        )
        self.assertEqual(len(rows), 1)
        self.assertFalse(rows[0]["Workload metrics available"])
        self.assertIsNone(rows[0]["Generated candidates (median)"])
        self.assertIsNone(rows[0]["Max/mean load ratio (median)"])
        self.assertEqual(rows[0]["Grid blocks (median)"], 320)

    def test_parallel_architecture_pairs_same_backend_and_thread_baselines(
        self,
    ) -> None:
        common = {
            "case_id": "parallel_v2_skewed",
            "semantic_group": "consecutive",
            "dataset": "skewed",
            "query": "reuse_star",
            "query_edges": 4,
            "k": 3,
            "suite": "parallel_v2_cpu_scaling",
            "threads": 32,
            "backend": "cpu",
            "device": "",
            "match_count": 17,
            "successful_runs": 5,
            "repeat_match_count_consistent": True,
        }
        summaries = [
            {
                **common,
                "algorithm": "cpu_prefix_t1",
                "algorithm_role": "parallel_v2_cpu_prefix_scheduler",
                "threads": 1,
                "runtime_s": 4.0,
                "source": {
                    "gopher_parallel_mode_effective_value": "roots-single-worker",
                },
            },
            {
                **common,
                "algorithm": "cpu_root_tasks_frozen_t32",
                "algorithm_role": "parallel_v2_cpu_root_tasks_frozen",
                "runtime_s": 2.0,
                "source": {},
            },
            {
                **common,
                "algorithm": "cpu_roots_v2_t32",
                "algorithm_role": "parallel_v2_cpu_current_roots",
                "runtime_s": 1.0,
                "source": {
                    "gopher_parallel_mode_effective_value": "roots",
                },
            },
            {
                **common,
                "algorithm": "cpu_prefix_t32",
                "algorithm_role": "parallel_v2_cpu_prefix_scheduler",
                "runtime_s": 0.5,
                "source": {
                    "gopher_parallel_mode_effective_value": "prefix",
                    "gopher_selectivity_mode_value": "intended",
                    "gopher_order_value": "paper",
                    "gopher_statistics_enabled_value": False,
                    "gopher_parallel_prefix_tasks_spawned_median": 64,
                    "gopher_parallel_prefix_split_attempts_median": 80,
                    "gopher_parallel_prefix_split_rejected_median": 16,
                    "gopher_parallel_queue_peak_median": 128,
                    "gopher_parallel_task_queue_capacity_median": 256,
                    "gopher_parallel_workers_with_tasks_median": 32,
                    "gopher_parallel_worker_dfs_states_max_median": 20,
                    "gopher_parallel_worker_dfs_states_total_median": 200,
                    "gopher_parallel_worker_dfs_states_cv_median": 0.1,
                },
            },
        ]
        rows = create_excel.build_parallel_architecture_rows(
            summaries,
            [
                {
                    "case_id": "parallel_v2_skewed",
                    "semantic_group": "consecutive",
                    "status": "pass",
                }
            ],
        )
        self.assertEqual(len(rows), 4)
        prefix = next(
            row
            for row in rows
            if row["Structure"] == "prefix" and row["Threads"] == 32
        )
        self.assertEqual(
            prefix["Baseline algorithm"],
            "cpu_root_tasks_frozen_t32",
        )
        self.assertAlmostEqual(
            prefix["Speedup vs same-backend old structure (×)"], 4.0
        )
        self.assertEqual(prefix["Current control algorithm"], "cpu_roots_v2_t32")
        self.assertAlmostEqual(prefix["Speedup vs current control (×)"], 2.0)
        self.assertEqual(prefix["Prefix tasks spawned"], 64)
        self.assertAlmostEqual(prefix["Prefix split rejection rate"], 0.2)
        self.assertAlmostEqual(prefix["Queue peak / capacity"], 0.5)
        self.assertAlmostEqual(prefix["Workers-with-tasks ratio"], 1.0)
        self.assertAlmostEqual(prefix["Worker DFS max / mean"], 3.2)
        self.assertEqual(prefix["Selectivity"], "intended")
        self.assertEqual(prefix["Order"], "paper")
        self.assertFalse(prefix["Statistics enabled"])
        self.assertAlmostEqual(prefix["Worker DFS-state CV"], 0.1)
        self.assertAlmostEqual(
            prefix["CPU structure speedup vs 1T (×)"], 8.0
        )
        self.assertAlmostEqual(prefix["CPU parallel efficiency"], 0.25)
        self.assertAlmostEqual(
            prefix["CPU Karp-Flatt fraction"], 3.0 / 31.0
        )

    def test_dimension_sweeps_hold_the_other_axis_constant(self) -> None:
        base = {
            "dataset": "wiki",
            "query": "path",
            "semantic_group": "consecutive",
            "algorithm": "ours",
            "runtime_s": 1.0,
            "match_count": 1,
            "successful_runs": 3,
        }
        summaries = [
            {
                **base,
                "case_id": f"paper_wiki_path{edges}_k{k}",
                "query_edges": edges,
                "k": k,
            }
            for edges, k in ((3, 3), (4, 3), (5, 3), (3, 4), (3, 5), (4, 4))
        ]
        query_rows = create_excel.build_dimension_rows(
            summaries, "query_edges"
        )
        self.assertEqual({row["k"] for row in query_rows}, {3})
        self.assertEqual(
            {row["Query edges"] for row in query_rows}, {3, 4, 5}
        )
        k_rows = create_excel.build_dimension_rows(summaries, "k")
        self.assertEqual({row["Query edges"] for row in k_rows}, {3})
        self.assertEqual({row["k"] for row in k_rows}, {3, 4, 5})

    def run_generator(self, fixture: Path, output: Path) -> dict[str, object]:
        completed = subprocess.run(
            [
                sys.executable,
                str(SCRIPT),
                "--input-dir",
                str(fixture),
                "--output",
                str(output),
                "--title",
                "Synthetic DSM Excel E2E Fixture",
            ],
            cwd=HERE.parent,
            text=True,
            capture_output=True,
            env=os.environ.copy(),
            check=False,
        )
        self.assertEqual(
            completed.returncode,
            0,
            msg=f"stdout:\n{completed.stdout}\nstderr:\n{completed.stderr}",
        )
        receipt = json.loads(completed.stdout)
        self.assertEqual(receipt["validation"], "passed")
        self.assertTrue(output.is_file())
        return receipt

    def test_full_fixture_builds_auditable_workbook(self) -> None:
        with tempfile.TemporaryDirectory(prefix="dsm_excel_test_") as temp:
            output = Path(temp) / "benchmark_fixture.xlsx"
            receipt = self.run_generator(FULL_FIXTURE, output)
            self.assertGreaterEqual(receipt["charts"], 8)
            self.assertGreaterEqual(receipt["formulas"], 10)

            validated = create_excel.validate_workbook(
                output,
                minimum_charts=8,
                required_chart_sheets=(
                    "Overview",
                    "Datasets",
                    "Speedup",
                    "Query Size",
                    "K Sensitivity",
                    "Gopher Ablation",
                    "CPU Scaling",
                    "CPU Single Thread",
                ),
            )
            self.assertEqual(validated["validation"], "passed")

            workbook = load_workbook(output, data_only=False)
            self.assertEqual(tuple(workbook.sheetnames), create_excel.SHEET_NAMES)

            for sheet in workbook.worksheets:
                self.assertIsNotNone(sheet.freeze_panes, sheet.title)
                if sheet.title != "Overview":
                    self.assertIsNotNone(sheet.auto_filter.ref, sheet.title)

            raw = workbook["Raw Runs"]
            raw_headers = {
                cell.value: cell.column for cell in raw[1] if cell.value is not None
            }
            self.assertIn("unknown_future_metric", raw_headers)
            statuses = {
                raw.cell(row=row, column=raw_headers["status"]).value
                for row in range(2, raw.max_row + 1)
            }
            self.assertIn("timeout", statuses)
            self.assertIn("error", statuses)

            summary = workbook["Summary"]
            summary_headers = {
                cell.value: cell.column
                for cell in summary[1]
                if cell.value is not None
            }
            self.assertIn("unknown_future_metric", summary_headers)
            unknown_values = [
                summary.cell(
                    row=row, column=summary_headers["unknown_future_metric"]
                ).value
                for row in range(2, summary.max_row + 1)
            ]
            self.assertIn("'=1+1", unknown_values)
            self.assertGreater(len(summary.conditional_formatting), 0)

            environment = workbook["Environment & GPU"]
            environment_rows = [
                tuple(cell.value for cell in row)
                for row in environment.iter_rows(min_row=2, max_col=3)
            ]
            self.assertIn(
                ("GPU / accelerator", "gpu.available", False),
                environment_rows,
            )

            speedup = workbook["Speedup"]
            speedup_headers = {
                cell.value: cell.column
                for cell in speedup[1]
                if cell.value is not None
            }
            repeat_row = None
            cardinality_rows = []
            for row in range(2, speedup.max_row + 1):
                case_id = speedup.cell(
                    row=row, column=speedup_headers["Case ID"]
                ).value
                algorithm = speedup.cell(
                    row=row, column=speedup_headers["Algorithm"]
                ).value
                if case_id == "bitcoin_repeat_path3_k3" and algorithm == "ours":
                    repeat_row = row
                if (
                    speedup.cell(
                        row=row, column=speedup_headers["Semantics"]
                    ).value
                    == "cardinality"
                ):
                    cardinality_rows.append(row)
            self.assertIsNotNone(repeat_row)
            self.assertEqual(
                speedup.cell(
                    row=repeat_row,
                    column=speedup_headers["Comparison status"],
                ).value,
                "not comparable",
            )
            self.assertIsNone(
                speedup.cell(
                    row=repeat_row, column=speedup_headers["Speedup (×)"]
                ).value
            )
            self.assertTrue(cardinality_rows)
            for row in cardinality_rows:
                self.assertEqual(
                    speedup.cell(
                        row=row, column=speedup_headers["Algorithm"]
                    ).value,
                    "original",
                )
                self.assertEqual(
                    speedup.cell(
                        row=row, column=speedup_headers["Baseline"]
                    ).value,
                    "original",
                )

            charts = [
                chart
                for sheet in workbook.worksheets
                for chart in sheet._charts
            ]
            self.assertTrue(charts)
            self.assertTrue(
                all(
                    chart.__class__.__name__ in {"BarChart", "LineChart"}
                    for chart in charts
                )
            )
            scaling = workbook["CPU Scaling"]
            scaling_headers = {
                cell.value: cell.column
                for cell in scaling[1]
                if cell.value is not None
            }
            self.assertIn("Scaling series", scaling_headers)
            self.assertGreaterEqual(len(scaling._charts), 1)

            validated_cli = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    "--validate-only",
                    str(output),
                ],
                cwd=HERE.parent,
                text=True,
                capture_output=True,
                check=False,
            )
            self.assertEqual(
                validated_cli.returncode,
                0,
                msg=validated_cli.stderr,
            )

    def test_sparse_fixture_keeps_empty_dimension_sheets_valid(self) -> None:
        with tempfile.TemporaryDirectory(prefix="dsm_excel_sparse_") as temp:
            output = Path(temp) / "sparse_fixture.xlsx"
            receipt = self.run_generator(SPARSE_FIXTURE, output)
            self.assertEqual(receipt["charts"], 1)

            workbook = load_workbook(output, data_only=False)
            self.assertEqual(len(workbook["Overview"]._charts), 1)
            for name in (
                "Datasets",
                "Speedup",
                "Query Size",
                "K Sensitivity",
                "Gopher Ablation",
                "CPU Scaling",
                "CPU Single Thread",
                "GPU Evaluation",
            ):
                self.assertEqual(len(workbook[name]._charts), 0, name)
                self.assertIsNotNone(workbook[name].freeze_panes)
                self.assertIsNotNone(workbook[name].auto_filter.ref)
            formulas = [
                cell.value
                for row in workbook["Overview"].iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ]
            self.assertTrue(formulas)
            self.assertFalse(any("#REF!" in formula for formula in formulas))


if __name__ == "__main__":
    unittest.main(verbosity=2)
